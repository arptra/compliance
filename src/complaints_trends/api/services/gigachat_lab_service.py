from __future__ import annotations

import concurrent.futures
import json
import math
import re
import shutil
import threading
import traceback
import uuid
from contextlib import suppress
from io import BytesIO
from datetime import date, datetime, timezone
from pathlib import Path
from typing import Any
from zipfile import ZipFile

import pandas as pd
from openpyxl.styles import PatternFill

from ...config import LLMConfig, ProjectConfig
from ...gigachat_api import build_gigachat_transport_client
from ...gigachat_mtls import SYSTEM_PROMPT
from ..schemas import (
    GigaChatAnnotatedExportRequest,
    GigaChatBackgroundTaskInputRow,
    GigaChatBackgroundTaskListResponse,
    GigaChatBackgroundTaskResultResponse,
    GigaChatBackgroundTaskRowRun,
    GigaChatBackgroundTaskStartRequest,
    GigaChatBackgroundTaskSummary,
    GigaChatFinalPromptRequest,
    GigaChatFinalPromptResponse,
    GigaChatLabRowRunRequest,
    GigaChatLabRowRunResponse,
    GigaChatRuleEvaluationRequest,
    GigaChatRuleEvaluationResponse,
    GigaChatRuleEvaluationRow,
    GigaChatRuleHit,
    GigaChatRulePack,
    GigaChatRulePackFilter,
    GigaChatReclassificationRuleImportItem,
    GigaChatReclassificationRulesImportResponse,
    GigaChatLabSettingField,
    GigaChatLabSettingOption,
    GigaChatLabSettingsResponse,
    GigaChatLabSettingsUpdateRequest,
    GigaChatLabSettingsVersionCreateRequest,
    GigaChatLabSettingsVersionResponse,
    GigaChatLabSettingsVersionsResponse,
    GigaChatLabSettingsVersionSummary,
    GigaChatLabSettingsVersionUpdateRequest,
    GigaChatWorkbookSelectSheetRequest,
    GigaChatWorkbookChunkedUploadCompleteResponse,
    GigaChatWorkbookChunkedUploadStartRequest,
    GigaChatWorkbookChunkedUploadStartResponse,
    GigaChatWorkbookChunkUploadResponse,
    GigaChatWorkbookRowsExportRequest,
    GigaChatWorkbookSheetDataResponse,
    GigaChatWorkbookSheetPreview,
    GigaChatWorkbookUploadTaskResponse,
    GigaChatWorkbookUploadResponse,
)

try:
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE as _OPENPYXL_ILLEGAL_RE
except Exception:  # pragma: no cover
    _OPENPYXL_ILLEGAL_RE = re.compile(r"[\000-\010]|[\013-\014]|[\016-\037]")


class _WorkbookUploadCancelled(Exception):
    pass


class _BackgroundTaskCancelled(Exception):
    pass


class GigaChatLabService:
    DEFAULT_SYSTEM_PROMPT = (
        "Ты обязан вернуть только валидный JSON без markdown, без пояснений вне JSON и без служебного текста. "
        "Размечай одно клиентское обращение за раз, используй все доступные поля строки и не выдумывай факты, категории или теги."
    )
    DEFAULT_USER_PROMPT_PREFIX = (
        "Разметь одну запись и верни JSON с полями: "
        "client_first_message, short_summary, is_complaint, complaint_category, complaint_subcategory, "
        "product_area, loan_product, severity, keywords, assigned_tags, local_tags, model_added_tags, "
        "model_rejected_tags, reclassified_topic, confirmed_rule_hits, rejected_rule_hits, tag_decisions, "
        "match_type, evidence, confidence, notes, evidence_columns. "
        "Сначала определи, является ли запись жалобой. Затем выбери основную категорию, при необходимости подкатегорию, "
        "назначь только разрешенные теги, при необходимости подтверди или отвергни подсказки rule-based движка, "
        "оцени severity и confidence, а в notes кратко объясни решение. "
        "Если категория неочевидна, верни пустую строку в complaint_category и объясни причину в notes."
    )
    LEGACY_CLASS_USER_PROMPT_PREFIX = (
        "Разметь одну запись и верни JSON с полями: "
        "client_first_message, short_summary, is_complaint, complaint_category, complaint_subcategory, "
        "product_area, loan_product, severity, keywords, assigned_tags, local_tags, model_added_tags, "
        "model_rejected_tags, reclassified_topic, confirmed_rule_hits, rejected_rule_hits, tag_decisions, "
        "match_type, evidence, confidence, notes, evidence_columns. "
        "Сначала определи, является ли запись жалобой. Затем выбери основную категорию, при необходимости подкатегорию, "
        "назначь только разрешенные теги, при необходимости подтверди или отвергни подсказки rule-based движка, "
        "оцени severity и confidence, а в notes кратко объясни решение. Возвращай только классы и теги из разрешенных списков. "
        "Если ни один разрешенный класс не подходит, верни пустую строку в complaint_category и объясни причину в notes."
    )
    DEFAULT_RECLASSIFICATION_PROMPT = ""
    EMPTY_RECLASSIFICATION_MARKERS = {
        "-",
        "—",
        "–",
        "нет",
        "нет изменений",
        "без изменений",
        "не менять",
        "оставить",
        "оставить как есть",
        "none",
        "null",
        "nil",
        "n/a",
        "na",
        "no change",
        "no_change",
        "same",
        "same topic",
        "same_topic",
    }
    DEFAULT_CONTEXT_NOTES = (
        "Источник данных: Excel/CSV с клиентскими обращениями. "
        "Одна строка таблицы = одно обращение для разметки. "
        "Используй одновременно текст клиента, комментарии, суммаризацию, продукт, тематику, канал и другие непустые поля. "
        "Если поля противоречат друг другу, приоритет у явного текста клиента, затем у суммаризации, затем у служебных атрибутов. "
        "Если обращение не выглядит как жалоба, верни is_complaint=false и не придумывай complaint_category."
    )
    DEFAULT_FALLBACK_ON_BLANK_KEYS = {
        "system_prompt",
        "user_prompt_prefix",
        "context_notes",
        "rule_pack_prompt_notes",
        "rule_pack_exclusion_notes",
        "reclassification_prompt_notes",
        "reclassification_prompt",
    }
    RULE_FILTER_FIELD_ALIASES: dict[str, list[str]] = {
        "Трайб": ["Во. Группа"],
        "драйвер": ["Во. Тематика"],
    }
    FIELD_DEFS: list[dict[str, Any]] = [
        {"key": "model", "label": "Model", "input_type": "text", "section": "Запрос"},
        {"key": "temperature", "label": "Temperature", "input_type": "number", "section": "Запрос", "help_text": "Насколько вариативным делать ответ модели."},
        {"key": "top_p", "label": "Top P", "input_type": "number", "section": "Запрос", "help_text": "Альтернативный sampling-порог по вероятности токенов."},
        {"key": "max_output_tokens", "label": "Max output tokens", "input_type": "number", "section": "Запрос", "help_text": "Максимальный размер ответа модели."},
        {"key": "system_prompt", "label": "System prompt", "input_type": "textarea", "section": "Промпты", "help_text": "Базовый system prompt для экспериментальной разметки."},
        {"key": "user_prompt_prefix", "label": "User prompt prefix", "input_type": "textarea", "section": "Промпты", "help_text": "Дополнительный текст перед пользовательским payload."},
        {"key": "context_notes", "label": "Context notes", "input_type": "textarea", "section": "Промпты", "help_text": "Текстовые инструкции про контекст, листы Excel и особенности эксперимента."},
        {"key": "rule_pack_prompt_notes", "label": "Rule packs", "input_type": "textarea", "section": "Разметка", "help_text": "Локальные rule-based пакеты: фильтры по полям, словари и действия до GigaChat."},
        {"key": "rule_pack_exclusion_notes", "label": "Rule pack exclusions", "input_type": "textarea", "section": "Разметка", "help_text": "JSON-массив code rule packs, которые остаются в локальной таблице, но не попадают в prompt/API GigaChat."},
        {"key": "reclassification_prompt_notes", "label": "Reclassification rules", "input_type": "textarea", "section": "Переклассификация", "help_text": "JSON-массив правил переклассификации: name, source_field, context_field/context_fields, prompt."},
    ]
    _background_lock = threading.Lock()
    _background_cancel_flags: dict[str, threading.Event] = {}
    _background_threads: dict[str, threading.Thread] = {}
    _upload_task_lock = threading.Lock()
    _upload_cancel_flags: dict[str, threading.Event] = {}
    _upload_task_threads: dict[str, threading.Thread] = {}

    def __init__(self, cfg: ProjectConfig, catalog_service: Any | None = None, lake_service: Any | None = None) -> None:
        self.cfg = cfg
        self.catalog = catalog_service
        self.lake = lake_service
        self.base_dir = Path(cfg.analysis.pattern_monitoring.interim_dir) / "gigachat_lab"
        self.uploads_dir = self.base_dir / "uploads"
        self.chunked_uploads_dir = self.base_dir / "chunked_uploads"
        self.background_dir = Path("data/background")
        self.versions_dir = Path("data/gigachat_lab/versions")
        self.settings_path = self.base_dir / "settings.json"
        self._transport_client_lock = threading.Lock()
        self._transport_clients: dict[tuple[Any, ...], Any] = {}
        self.base_dir.mkdir(parents=True, exist_ok=True)
        self.uploads_dir.mkdir(parents=True, exist_ok=True)
        self.chunked_uploads_dir.mkdir(parents=True, exist_ok=True)
        self.background_dir.mkdir(parents=True, exist_ok=True)
        self.versions_dir.mkdir(parents=True, exist_ok=True)
        self._seed_catalog_settings_versions()

    @staticmethod
    def _now() -> datetime:
        return datetime.now(timezone.utc)

    def _default_rule_packs(self) -> str:
        items = [
            {
                "code": "DRA",
                "description": "Проставляет тег DRA/ДРПА по словам про смерть, наследство, каникулы, реструктуризацию, приставов, СВО, суд, исполнительное производство, военный контур и банкротство.",
                "enabled": True,
                "type": "assign_tag",
                "source_fields": ["Во. Описание", "Обр. Результат суммаризации диалога"],
                "keywords": [
                    "умер",
                    "погиб",
                    "смерт",
                    "гибел",
                    "наследни",
                    "наследств",
                    "каникул",
                    "реструктуриз",
                    "пристав",
                    "участник СВО",
                    "участника СВО",
                    "участником СВО",
                    "на СВО",
                    "судебное решение",
                    "по решению суда",
                    "исполнительное производство",
                    "военн",
                    "банкрот",
                ],
                "filters": [],
                "target_tag": "DRA",
                "target_topic": None,
            },
            {
                "code": "IPOTEKA",
                "description": "Проставляет тег ИПОТЕКА по обращениям про ипотеку, жилищный кредит, недвижимость в залоге, закладную, эскроу, обременение и ипотечное рефинансирование.",
                "enabled": True,
                "type": "assign_tag",
                "source_fields": ["Во. Описание", "Обр. Результат суммаризации диалога"],
                "keywords": [
                    "ипотек",
                    "жилищн",
                    "недвижим",
                    "квартир",
                    "дом в залог",
                    "залог недвиж",
                    "закладн",
                    "эскроу",
                    "обременен",
                    "созаемщик",
                    "созаемщ",
                    "рефинансир*ипот",
                    "рефинансирован*ипот",
                    "материнск*капитал",
                    "первоначальн*взнос",
                ],
                "filters": [],
                "target_tag": "ИПОТЕКА",
                "target_topic": None,
            },
            {
                "code": "EDU_RECLASS_TRANCH",
                "description": "Переклассифицирует обращения в тему про очередной транш по образовательному кредиту.",
                "enabled": True,
                "type": "reclass_topic",
                "source_fields": ["Во. Описание", "Обр. Результат суммаризации диалога"],
                "keywords": ["транш", "семестр"],
                "filters": [
                    {"field": "Трайб", "op": "eq", "value": "ПОТРЕБИТЕЛЬСКИЕ КРЕДИТЫ"},
                    {"field": "драйвер", "op": "ne", "value": "ОБРАЗОВАТЕЛЬНЫЙ КРЕДИТ"},
                ],
                "target_tag": None,
                "target_topic": "Проблема с выдачей очередного транша по Образовательному кредиту",
            },
            {
                "code": "EDU_RECLASS_APPLICATION",
                "description": "Переклассифицирует обращения в тему про зачисление средств, оформление или рассмотрение заявки по образовательному кредиту.",
                "enabled": True,
                "type": "reclass_topic",
                "source_fields": ["Во. Описание", "Обр. Результат суммаризации диалога"],
                "keywords": [
                    "Образовательн",
                    "Вуз",
                    "Кредит на образ",
                    "Оплатить обучение",
                    "Период*обучения",
                    "Отчисл",
                ],
                "filters": [
                    {"field": "Трайб", "op": "eq", "value": "ПОТРЕБИТЕЛЬСКИЕ КРЕДИТЫ"},
                    {"field": "драйвер", "op": "ne", "value": "ОБРАЗОВАТЕЛЬНЫЙ КРЕДИТ"},
                ],
                "target_tag": None,
                "target_topic": "Проблемы с зачислением средств/ оформлением-рассмотрением заявки",
            },
        ]
        return json.dumps(items, ensure_ascii=False, indent=2)

    def _default_values(self) -> dict[str, Any]:
        llm = self.cfg.llm.model_dump()
        default_system_prompt = str(llm.get("system_prompt", "") or "").strip()
        if not default_system_prompt or default_system_prompt == LLMConfig.model_fields["system_prompt"].default:
            default_system_prompt = self.DEFAULT_SYSTEM_PROMPT

        default_user_prompt_prefix = str(llm.get("user_prompt_prefix", "") or "").strip()
        if not default_user_prompt_prefix or default_user_prompt_prefix == self.LEGACY_CLASS_USER_PROMPT_PREFIX:
            default_user_prompt_prefix = self.DEFAULT_USER_PROMPT_PREFIX
        default_context_notes = str(llm.get("context_notes", "") or "").strip() or self.DEFAULT_CONTEXT_NOTES
        default_rule_pack_notes = str(llm.get("rule_pack_prompt_notes", "") or "").strip() or self._default_rule_packs()
        default_reclassification_notes = str(llm.get("reclassification_prompt_notes", "") or "").strip()
        default_reclassification_prompt = str(llm.get("reclassification_prompt", "") or "").strip() or self.DEFAULT_RECLASSIFICATION_PROMPT
        return {
            "model": llm.get("model", "GigaChat"),
            "temperature": llm.get("temperature", 0.2),
            "top_p": llm.get("top_p", 0.95),
            "max_output_tokens": llm.get("max_output_tokens", 2048),
            "system_prompt": default_system_prompt,
            "user_prompt_prefix": default_user_prompt_prefix,
            "context_notes": default_context_notes,
            "rule_pack_prompt_notes": default_rule_pack_notes,
            "rule_pack_exclusion_notes": "[]",
            "reclassification_prompt_notes": default_reclassification_notes,
            "reclassification_source_field": str(llm.get("reclassification_source_field", "") or "").strip(),
            "reclassification_context_field": str(llm.get("reclassification_context_field", "") or "").strip(),
            "reclassification_prompt": default_reclassification_prompt,
        }

    def _load_saved_settings(self) -> dict[str, Any]:
        if not self.settings_path.exists():
            return {"saved_at": None, "values": {}}
        try:
            payload = json.loads(self.settings_path.read_text(encoding="utf-8"))
            if not isinstance(payload, dict):
                return {"saved_at": None, "values": {}}
            values = payload.get("values", {})
            if not isinstance(values, dict):
                values = {}
            return {"saved_at": payload.get("saved_at"), "values": values}
        except Exception:
            return {"saved_at": None, "values": {}}

    def _effective_values(self) -> tuple[dict[str, Any], str | None]:
        defaults = self._default_values()
        saved = self._load_saved_settings()
        values = dict(defaults)
        for key, value in dict(saved.get("values", {})).items():
            if key not in defaults:
                continue
            if key in self.DEFAULT_FALLBACK_ON_BLANK_KEYS and (value is None or (isinstance(value, str) and not value.strip())):
                continue
            if key == "user_prompt_prefix" and str(value or "").strip() == self.LEGACY_CLASS_USER_PROMPT_PREFIX:
                continue
            values[key] = value
        return values, saved.get("saved_at")

    def _fields_from_values(self, values: dict[str, Any]) -> list[GigaChatLabSettingField]:
        fields: list[GigaChatLabSettingField] = []
        for meta in self.FIELD_DEFS:
            options = [GigaChatLabSettingOption(value=v, label=l) for v, l in meta.get("options", [])]
            fields.append(
                GigaChatLabSettingField(
                    key=meta["key"],
                    label=meta["label"],
                    input_type=meta["input_type"],
                    section=meta["section"],
                    help_text=meta.get("help_text"),
                    value=values.get(meta["key"]),
                    options=options,
                )
            )
        return fields

    @staticmethod
    def _version_id_from_title(title: str) -> str:
        raw = re.sub(r"[^a-zA-Z0-9_-]+", "_", title.strip().lower()).strip("_")
        return raw or f"version_{uuid.uuid4().hex[:8]}"

    def _version_path(self, version_id: str) -> Path:
        safe = self._version_id_from_title(version_id)
        return self.versions_dir / f"{safe}.json"

    @staticmethod
    def _user_context(user: dict[str, Any] | None = None, *, user_id: str | None = None, workspace_id: str | None = None) -> dict[str, str]:
        if user:
            return {
                "user_id": str(user.get("id") or user_id or "anonymous"),
                "workspace_id": str(user.get("workspace_id") or workspace_id or "default"),
                "display_name": str(user.get("display_name") or user.get("email") or user_id or "anonymous"),
            }
        return {
            "user_id": str(user_id or "anonymous"),
            "workspace_id": str(workspace_id or "default"),
            "display_name": str(user_id or "anonymous"),
        }

    def _default_version_payload(self) -> dict[str, Any]:
        values, saved_at = self._effective_values()
        now = saved_at or self._now().isoformat()
        return {
            "version_id": "default",
            "title": "Default",
            "description": "Базовая версия из дефолтных и legacy-настроек Lab.",
            "status": "release",
            "created_by": "system",
            "owner_user_id": "system",
            "workspace_id": "default",
            "visibility": "public",
            "created_at": now,
            "updated_at": saved_at,
            "base_version_id": None,
            "values": values,
        }

    def _seed_catalog_settings_versions(self) -> None:
        if not self.catalog:
            return
        default_payload = self._default_version_payload()
        if not self.catalog.lab_settings_version_exists("default"):
            self.catalog.save_lab_settings_version(default_payload)
        for path in sorted(self.versions_dir.glob("*.json")):
            with suppress(Exception):
                payload = json.loads(path.read_text(encoding="utf-8"))
                if not isinstance(payload, dict):
                    continue
                version_id = str(payload.get("version_id") or path.stem)
                if self.catalog.lab_settings_version_exists(version_id):
                    continue
                payload["version_id"] = version_id
                payload.setdefault("workspace_id", "default")
                payload.setdefault("owner_user_id", "system")
                payload.setdefault("visibility", "public")
                payload.setdefault("created_by", payload.get("created_by") or "system")
                payload.setdefault("updated_by", payload.get("updated_by") or payload.get("created_by") or "")
                payload.setdefault("metadata", {"migrated_from": str(path)})
                self.catalog.save_lab_settings_version(payload)

    def _read_version_payload(self, version_id: str, *, user: dict[str, Any] | None = None, user_id: str | None = None, workspace_id: str | None = None) -> dict[str, Any]:
        if self.catalog:
            context = self._user_context(user, user_id=user_id, workspace_id=workspace_id)
            payload = self.catalog.get_lab_settings_version(
                version_id,
                user_id=context["user_id"],
                workspace_id=context["workspace_id"],
            )
            if payload:
                return payload
            if version_id == "default":
                return self._default_version_payload()
            raise FileNotFoundError(f"Settings version not found: {version_id}")
        if version_id == "default" and not self._version_path(version_id).exists():
            return self._default_version_payload()
        path = self._version_path(version_id)
        if not path.exists():
            raise FileNotFoundError(f"Settings version not found: {version_id}")
        payload = json.loads(path.read_text(encoding="utf-8"))
        if not isinstance(payload, dict):
            raise ValueError(f"Settings version is not a JSON object: {version_id}")
        return payload

    def _write_version_payload(self, payload: dict[str, Any]) -> None:
        version_id = str(payload.get("version_id") or "")
        if not version_id:
            raise ValueError("version_id is required")
        if self.catalog:
            self.catalog.save_lab_settings_version(payload)
            return
        self._version_path(version_id).write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")

    def _version_summary_from_payload(self, payload: dict[str, Any], *, is_default: bool = False, user_id: str | None = None) -> GigaChatLabSettingsVersionSummary:
        version_id = str(payload.get("version_id") or "default")
        path = self._version_path(version_id)
        owner_user_id = str(payload.get("owner_user_id") or "")
        resolved_is_default = is_default or version_id == "default" or bool(payload.get("is_default"))
        return GigaChatLabSettingsVersionSummary(
            version_id=version_id,
            title=str(payload.get("title") or version_id),
            description=str(payload.get("description") or ""),
            status=payload.get("status") or "draft",
            visibility=payload.get("visibility") or "private",
            created_by=str(payload.get("created_by") or ""),
            owner_user_id=owner_user_id,
            created_at=datetime.fromisoformat(payload["created_at"]) if payload.get("created_at") else None,
            updated_at=datetime.fromisoformat(payload["updated_at"]) if payload.get("updated_at") else None,
            base_version_id=payload.get("base_version_id"),
            path=str(payload.get("path") or (str(path) if path.exists() else None)) if not self.catalog else "sqlite:data/app.sqlite",
            is_default=resolved_is_default,
            can_edit=bool(payload.get("can_edit")) if "can_edit" in payload else (bool(user_id) and owner_user_id == user_id and not resolved_is_default),
        )

    def _values_from_version_payload(self, payload: dict[str, Any]) -> dict[str, Any]:
        defaults = self._default_values()
        values = dict(defaults)
        raw_values = payload.get("values", {})
        if isinstance(raw_values, dict):
            for key, value in raw_values.items():
                if key in defaults:
                    values[key] = self._coerce_setting_value(key, value)
        return values

    def list_settings_versions(self, *, user: dict[str, Any] | None = None) -> GigaChatLabSettingsVersionsResponse:
        if self.catalog:
            context = self._user_context(user)
            versions = [
                self._version_summary_from_payload(payload, user_id=context["user_id"])
                for payload in self.catalog.list_lab_settings_versions(user_id=context["user_id"], workspace_id=context["workspace_id"])
            ]
            if not versions:
                versions.append(self._version_summary_from_payload(self._default_version_payload(), is_default=True, user_id=context["user_id"]))
            return GigaChatLabSettingsVersionsResponse(versions=versions)
        versions: list[GigaChatLabSettingsVersionSummary] = []
        for path in sorted(self.versions_dir.glob("*.json")):
            with suppress(Exception):
                payload = json.loads(path.read_text(encoding="utf-8"))
                versions.append(self._version_summary_from_payload(payload, is_default=payload.get("version_id") == "default"))
        if not versions:
            versions.append(self._version_summary_from_payload(self._default_version_payload(), is_default=True))
        versions.sort(key=lambda item: (not item.is_default, item.title.lower()))
        return GigaChatLabSettingsVersionsResponse(versions=versions)

    def get_settings_version(self, version_id: str, *, user: dict[str, Any] | None = None) -> GigaChatLabSettingsVersionResponse:
        context = self._user_context(user)
        payload = self._read_version_payload(version_id, user=user)
        values = self._values_from_version_payload(payload)
        return GigaChatLabSettingsVersionResponse(
            version=self._version_summary_from_payload(payload, is_default=version_id == "default", user_id=context["user_id"]),
            fields=self._fields_from_values(values),
            values=values,
        )

    def create_settings_version(self, req: GigaChatLabSettingsVersionCreateRequest, *, user: dict[str, Any] | None = None) -> GigaChatLabSettingsVersionResponse:
        context = self._user_context(user)
        base_id = self._version_id_from_title(req.version_id or req.title)
        version_id = base_id
        if self.catalog:
            suffix = 2
            while self.catalog.lab_settings_version_exists(version_id):
                version_id = f"{base_id}_{suffix}"
                suffix += 1
        elif self._version_path(version_id).exists():
            raise ValueError(f"Settings version already exists: {version_id}")
        base_payload = self._read_version_payload(req.base_version_id or "default", user=user)
        now = self._now().isoformat()
        created_by = req.created_by.strip() or context["display_name"]
        payload = {
            "version_id": version_id,
            "workspace_id": context["workspace_id"],
            "owner_user_id": context["user_id"],
            "title": req.title.strip() or version_id,
            "description": req.description,
            "status": req.status,
            "visibility": req.visibility,
            "created_by": created_by,
            "updated_by": created_by,
            "created_at": now,
            "updated_at": now,
            "base_version_id": req.base_version_id or "default",
            "values": self._values_from_version_payload(base_payload),
        }
        self._write_version_payload(payload)
        return self.get_settings_version(version_id, user=user)

    def save_settings_version(self, version_id: str, req: GigaChatLabSettingsVersionUpdateRequest, *, user: dict[str, Any] | None = None) -> GigaChatLabSettingsVersionResponse:
        context = self._user_context(user)
        if version_id == "default" and not self._version_path(version_id).exists():
            payload = self._default_version_payload()
            payload["created_at"] = self._now().isoformat()
        else:
            payload = self._read_version_payload(version_id, user=user)
        if self.catalog:
            is_owner = str(payload.get("owner_user_id") or "") == context["user_id"]
            if version_id == "default" or not is_owner:
                raise PermissionError("Only the owner can update this settings version.")
        if req.title is not None:
            payload["title"] = req.title
        if req.description is not None:
            payload["description"] = req.description
        if req.status is not None:
            payload["status"] = req.status
        if req.visibility is not None:
            payload["visibility"] = req.visibility
        values = self._values_from_version_payload(payload)
        known_setting_keys = {str(field.get("key") or "") for field in self.FIELD_DEFS}
        for key, value in req.values.items():
            if key in values or key in known_setting_keys:
                values[key] = self._coerce_setting_value(key, value)
        payload["values"] = values
        payload["updated_by"] = req.updated_by or context["display_name"] or payload.get("updated_by") or payload.get("created_by") or ""
        payload["updated_at"] = self._now().isoformat()
        self._write_version_payload(payload)
        return self.get_settings_version(str(payload["version_id"]), user=user)

    def delete_settings_version(self, version_id: str, *, user: dict[str, Any] | None = None) -> GigaChatLabSettingsVersionsResponse:
        context = self._user_context(user)
        if version_id == "default":
            raise PermissionError("Default settings version cannot be deleted.")
        payload = self._read_version_payload(version_id, user=user)
        if self.catalog:
            is_owner = str(payload.get("owner_user_id") or "") == context["user_id"]
            if not is_owner:
                raise PermissionError("Only the owner can delete this settings version.")
            deleted = self.catalog.delete_lab_settings_version(
                version_id,
                user_id=context["user_id"],
                workspace_id=context["workspace_id"],
            )
            if not deleted:
                raise FileNotFoundError(f"Settings version not found: {version_id}")
            return self.list_settings_versions(user=user)

        path = self._version_path(version_id)
        if not path.exists():
            raise FileNotFoundError(f"Settings version not found: {version_id}")
        path.unlink()
        return self.list_settings_versions(user=user)

    def export_rule_packs_workbook(self, version_id: str, *, user: dict[str, Any] | None = None) -> tuple[str, bytes]:
        payload = self._read_version_payload(version_id, user=user)
        values = self._values_from_version_payload(payload)
        rule_packs = self._parse_rule_pack_items(values.get("rule_pack_prompt_notes", ""))
        columns = [
            "Название правила",
            "Поле поиска",
            "Ключевые слова",
            "Активно",
            "Тип действия",
            "Target tag",
            "Target topic",
            "Описание",
            "Фильтры JSON",
        ]
        rows = []
        for rule_pack in rule_packs:
            rows.append(
                {
                    "Название правила": rule_pack.code,
                    "Поле поиска": "\n".join(rule_pack.source_fields),
                    "Ключевые слова": "\n".join(rule_pack.keywords),
                    "Активно": "yes" if rule_pack.enabled else "no",
                    "Тип действия": rule_pack.type,
                    "Target tag": rule_pack.target_tag or "",
                    "Target topic": rule_pack.target_topic or "",
                    "Описание": rule_pack.description or "",
                    "Фильтры JSON": json.dumps([item.model_dump() for item in rule_pack.filters], ensure_ascii=False),
                }
            )

        df = pd.DataFrame(rows, columns=columns)
        guide = pd.DataFrame(
            [
                {"Поле": "Название тега", "Описание": "Код/название rule pack или тега, например DRA или IPOTEKA. Также принимается колонка Название правила."},
                {"Поле": "Колонка", "Описание": "Одна или несколько колонок, где ищутся ключевые слова. Также принимается колонка Поле поиска."},
                {"Поле": "Ключевые слова", "Описание": "Один или несколько ключей. Можно писать через перенос строки или запятую, * работает как wildcard."},
            ]
        )
        buffer = BytesIO()
        with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
            df.to_excel(writer, index=False, sheet_name="rules")
            guide.to_excel(writer, index=False, sheet_name="README")
            worksheet = writer.sheets["rules"]
            for index, width in enumerate([28, 36, 48, 12, 18, 22, 42, 56, 42], start=1):
                worksheet.column_dimensions[chr(64 + index)].width = width

        title = str(payload.get("title") or version_id or "rules")
        safe_title = re.sub(r"[^a-zA-Z0-9а-яА-ЯёЁ_-]+", "_", title).strip("_") or "rules"
        return f"{safe_title}_rule_packs.xlsx", buffer.getvalue()

    def import_rule_packs_workbook(
        self,
        version_id: str,
        filename: str,
        content: bytes,
        *,
        user: dict[str, Any] | None = None,
    ) -> GigaChatLabSettingsVersionResponse:
        context = self._user_context(user)
        suffix = Path(filename or "rules.xlsx").suffix.lower()
        if suffix == ".csv":
            df = pd.read_csv(BytesIO(content), dtype=str, keep_default_na=False)
        elif suffix in {".xlsx", ".xls", ".xlsm"}:
            df = pd.read_excel(BytesIO(content), sheet_name=0, dtype=str, keep_default_na=False)
        else:
            raise ValueError("Загрузите Excel-файл с правилами: .xlsx, .xls, .xlsm или .csv.")

        df = self._drop_empty_unnamed_columns(df).fillna("")
        column_map = {self._normalize_rule_pack_import_header(column): column for column in df.columns}

        def find_column(*aliases: str) -> str | None:
            for alias in aliases:
                column = column_map.get(self._normalize_rule_pack_import_header(alias))
                if column is not None:
                    return str(column)
            return None

        name_column = find_column(
            "Название правила",
            "Название тега",
            "Rule name",
            "Rule code",
            "Tag name",
            "Code",
            "Name",
            "Правило",
            "Тег",
        )
        fields_column = find_column("Поле поиска", "Поля поиска", "Source field", "Source fields", "Field", "Fields", "Колонка", "Колонки")
        keywords_column = find_column("Ключевые слова", "Keywords", "Keyword", "Ключи", "Ключевое слово")
        if not name_column or not fields_column or not keywords_column:
            raise ValueError("В Excel должны быть колонки: Название правила, Поле поиска, Ключевые слова.")

        enabled_column = find_column("Активно", "Enabled", "Active")
        type_column = find_column("Тип действия", "Type", "Action type")
        target_tag_column = find_column("Target tag", "Тег", "Tag")
        target_topic_column = find_column("Target topic", "Topic", "Тема")
        description_column = find_column("Описание", "Description")
        filters_column = find_column("Фильтры JSON", "Filters JSON", "Filters")

        rule_packs: list[GigaChatRulePack] = []
        for row_index, row in df.iterrows():
            code = self._cell_text(row.get(name_column))
            source_fields = self._split_rule_list(row.get(fields_column))
            keywords = self._split_keyword_list(row.get(keywords_column))
            if not code and not source_fields and not keywords:
                continue
            if not code:
                raise ValueError(f"Строка {row_index + 2}: заполните название правила.")
            if not source_fields:
                raise ValueError(f"Строка {row_index + 2}: заполните поле поиска.")
            if not keywords:
                raise ValueError(f"Строка {row_index + 2}: заполните ключевые слова.")

            target_tag = self._cell_text(row.get(target_tag_column)) if target_tag_column else ""
            target_topic = self._cell_text(row.get(target_topic_column)) if target_topic_column else ""
            rule_type = self._cell_text(row.get(type_column)).lower() if type_column else ""
            if rule_type not in {"assign_tag", "reclass_topic"}:
                rule_type = "reclass_topic" if target_topic else "assign_tag"

            filters = self._parse_rule_pack_import_filters(row.get(filters_column) if filters_column else "")
            rule_packs.append(
                GigaChatRulePack(
                    code=code,
                    description=self._cell_text(row.get(description_column)) if description_column else "",
                    enabled=self._parse_rule_pack_import_bool(row.get(enabled_column), default=True) if enabled_column else True,
                    type=rule_type,
                    source_fields=source_fields,
                    keywords=keywords,
                    filters=filters,
                    target_tag=(target_tag or code) if rule_type == "assign_tag" else None,
                    target_topic=target_topic if rule_type == "reclass_topic" else None,
                )
            )

        if not rule_packs:
            raise ValueError("В файле не найдено ни одного правила.")

        serialized = json.dumps([item.model_dump() for item in rule_packs], ensure_ascii=False, indent=2)
        req = GigaChatLabSettingsVersionUpdateRequest(
            updated_by=context["display_name"],
            values={"rule_pack_prompt_notes": serialized},
        )
        return self.save_settings_version(version_id, req, user=user)

    def import_reclassification_rules_workbook(
        self,
        filename: str,
        content: bytes,
    ) -> GigaChatReclassificationRulesImportResponse:
        suffix = Path(filename or "reclassification_rules.xlsx").suffix.lower()
        if suffix == ".csv":
            df = pd.read_csv(BytesIO(content), dtype=str, keep_default_na=False)
        elif suffix in {".xlsx", ".xls", ".xlsm"}:
            df = pd.read_excel(BytesIO(content), sheet_name=0, dtype=str, keep_default_na=False)
        else:
            raise ValueError("Загрузите Excel-файл с правилами переклассификации: .xlsx, .xls, .xlsm или .csv.")

        df = self._drop_empty_unnamed_columns(df).fillna("")
        column_map = {self._normalize_rule_pack_import_header(column): column for column in df.columns}

        def find_column(*aliases: str) -> str | None:
            for alias in aliases:
                column = column_map.get(self._normalize_rule_pack_import_header(alias))
                if column is not None:
                    return str(column)
            return None

        name_column = find_column("name", "название", "название переклассификации", "topic", "тема")
        source_column = find_column("src_field", "source_field", "source field", "поле исходной темы", "исходное поле")
        context_column = find_column("context_field", "context field", "поле контекста", "контекст")
        prompt_column = find_column("prompt_field", "prompt", "description", "описание", "описание переклассификации")
        if not name_column or not source_column or not context_column or not prompt_column:
            raise ValueError("В Excel должны быть колонки: name, src_field, context_field, prompt_field. В context_field можно перечислить несколько колонок через запятую или перенос строки.")

        rules: list[GigaChatReclassificationRuleImportItem] = []
        for row_index, row in df.iterrows():
            name = self._cell_text(row.get(name_column))
            source_field = self._cell_text(row.get(source_column))
            context_fields = self._split_rule_list(row.get(context_column))
            context_field = context_fields[0] if context_fields else ""
            prompt = self._cell_text(row.get(prompt_column))
            if not name and not source_field and not context_fields and not prompt:
                continue
            if not name:
                raise ValueError(f"Строка {row_index + 2}: заполните name.")
            if not source_field and not context_fields:
                raise ValueError(f"Строка {row_index + 2}: заполните src_field или context_field.")
            if not prompt:
                raise ValueError(f"Строка {row_index + 2}: заполните prompt_field.")
            rules.append(
                GigaChatReclassificationRuleImportItem(
                    name=name,
                    source_field=source_field,
                    context_field=context_field,
                    context_fields=context_fields,
                    prompt=prompt,
                )
            )

        if not rules:
            raise ValueError("В файле не найдено ни одного правила переклассификации.")
        return GigaChatReclassificationRulesImportResponse(
            filename=Path(filename or "reclassification_rules.xlsx").name,
            imported_count=len(rules),
            rules=rules,
        )

    def _coerce_setting_value(self, key: str, value: Any) -> Any:
        meta = next((x for x in self.FIELD_DEFS if x["key"] == key), None)
        if meta is None:
            return value
        input_type = meta["input_type"]
        if input_type == "boolean":
            if isinstance(value, str):
                return value.strip().lower() in {"1", "true", "yes", "y", "on"}
            return bool(value)
        if input_type == "number":
            try:
                return float(value) if "." in str(value) and key in {"temperature", "top_p"} else int(value)
            except Exception:
                return 0.0 if key in {"temperature", "top_p"} else 0
        if value is None:
            return ""
        return str(value)

    def get_settings(self) -> GigaChatLabSettingsResponse:
        values, saved_at_raw = self._effective_values()
        saved_at = None
        if isinstance(saved_at_raw, str) and saved_at_raw.strip():
            try:
                saved_at = datetime.fromisoformat(saved_at_raw)
            except Exception:
                saved_at = None
        return GigaChatLabSettingsResponse(
            title="GigaChat Lab Settings",
            fields=self._fields_from_values(values),
            saved_at=saved_at,
        )

    def save_settings(self, req: GigaChatLabSettingsUpdateRequest) -> GigaChatLabSettingsResponse:
        allowed = {meta["key"] for meta in self.FIELD_DEFS}
        saved_values = dict(self._load_saved_settings().get("values", {}))
        for key, value in req.values.items():
            if key not in allowed:
                continue
            saved_values[key] = self._coerce_setting_value(key, value)
        payload = {
            "saved_at": self._now().isoformat(),
            "values": saved_values,
        }
        self.settings_path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
        return self.get_settings()

    def apply_llm_overrides(self, llm_cfg: LLMConfig) -> LLMConfig:
        values, _ = self._effective_values()
        fields = llm_cfg.__class__.model_fields
        for key, value in values.items():
            if key in fields:
                setattr(llm_cfg, key, value)
        return llm_cfg

    @staticmethod
    def _transport_client_cache_key(llm_cfg: LLMConfig, transport: str | None) -> tuple[Any, ...]:
        selected = str(transport or llm_cfg.mode or "mtls").strip().lower()
        return (
            selected,
            llm_cfg.base_url,
            llm_cfg.oauth_url,
            llm_cfg.oauth_scope,
            llm_cfg.authorization_key_file,
            llm_cfg.ca_bundle_file,
            llm_cfg.cert_file,
            llm_cfg.key_file,
            llm_cfg.verify_ssl_certs,
            llm_cfg.key_file_password_env,
        )

    def _get_transport_client(self, llm_cfg: LLMConfig, transport: str | None) -> Any:
        key = self._transport_client_cache_key(llm_cfg, transport)
        with self._transport_client_lock:
            client = self._transport_clients.get(key)
            if client is None:
                client = build_gigachat_transport_client(llm_cfg, transport=transport)
                self._transport_clients[key] = client
            return client

    def build_final_prompt(self, req: GigaChatFinalPromptRequest) -> GigaChatFinalPromptResponse:
        values = dict(self._effective_values()[0])
        for key, value in req.values.items():
            values[key] = self._coerce_setting_value(key, value)

        columns = [str(column) for column in req.columns if str(column).strip()]
        payload = self._compose_final_payload(values, columns)
        generated_at = self._now()

        return GigaChatFinalPromptResponse(
            generated_at=generated_at,
            source_columns=columns,
            payload=payload,
            saved=False,
            saved_path=None,
        )

    def run_row_prompt(self, req: GigaChatLabRowRunRequest) -> GigaChatLabRowRunResponse:
        values = dict(self._effective_values()[0])
        for key, value in req.values.items():
            values[key] = self._coerce_setting_value(key, value)

        columns = [str(column) for column in req.columns if str(column).strip()]
        base_payload = (
            req.payload_override
            if req.payload_override and not req.reclassification_only
            else (
                self._compose_reclassification_payload(values, columns)
                if req.reclassification_only
                else self._compose_final_payload(values, columns)
            )
        )
        rule_evaluation: GigaChatRuleEvaluationRow | None = None
        prompt_rule_evaluation: GigaChatRuleEvaluationRow | None = None
        if not req.reclassification_only:
            rule_packs = self._parse_rule_pack_items(values.get("rule_pack_prompt_notes", ""))
            rule_evaluation = self._evaluate_rule_hits_for_row(rule_packs, req.row, 0)
            prompt_rule_evaluation = self._filter_rule_evaluation_for_prompt(
                rule_evaluation,
                self._parse_rule_pack_exclusions(values.get("rule_pack_exclusion_notes", "")),
            )
        rendered_payload = self._render_payload_for_row(base_payload, req.row)
        if prompt_rule_evaluation is not None:
            rendered_payload = self._inject_row_rule_context(rendered_payload, prompt_rule_evaluation)

        llm_cfg = self.cfg.llm.model_copy(deep=True)
        llm_cfg = self.apply_llm_overrides(llm_cfg)
        llm_fields = llm_cfg.__class__.model_fields
        for key, value in values.items():
            if key in llm_fields:
                setattr(llm_cfg, key, value)
        llm_cfg.mode = req.transport

        client = self._get_transport_client(llm_cfg, req.transport)
        token_count = None
        if req.count_tokens and hasattr(client, "count_tokens"):
            token_count = client.count_tokens(model=llm_cfg.model, input_text=self._payload_to_count_input(rendered_payload))
        raw_content = client.chat(rendered_payload).choices[0].message.content
        parsed_payload: dict[str, Any] | list[Any] | None = None
        parse_ok = False
        try:
            parsed_payload = json.loads(raw_content)
            parse_ok = True
        except Exception:
            parsed_payload = None
        if req.reclassification_only:
            parsed_payload = self._sanitize_reclassification_response(parsed_payload, values)

        return GigaChatLabRowRunResponse(
            transport=req.transport,
            request_payload=rendered_payload,
            response_raw=raw_content,
            response_json=parsed_payload,
            parse_ok=parse_ok,
            request_token_count=token_count,
            rule_evaluation=rule_evaluation,
        )

    @classmethod
    def _values_without_reclassification(cls, values: dict[str, Any]) -> dict[str, Any]:
        next_values = dict(values)
        next_values["reclassification_prompt_notes"] = "[]"
        next_values["reclassification_source_field"] = ""
        next_values["reclassification_context_field"] = ""
        next_values["reclassification_prompt"] = cls.DEFAULT_RECLASSIFICATION_PROMPT
        return next_values

    @staticmethod
    def _coerce_async_workers(value: Any) -> int:
        try:
            workers = int(value)
        except (TypeError, ValueError):
            workers = 1
        return max(1, min(32, workers))

    def list_background_tasks(self) -> GigaChatBackgroundTaskListResponse:
        tasks: list[GigaChatBackgroundTaskSummary] = []
        for meta_path in sorted(self.background_dir.glob("*/task.json"), reverse=True):
            with suppress(Exception):
                tasks.append(GigaChatBackgroundTaskSummary(**json.loads(meta_path.read_text(encoding="utf-8"))))
        tasks.sort(key=lambda item: item.created_at, reverse=True)
        return GigaChatBackgroundTaskListResponse(tasks=tasks)

    def start_background_labeling(self, req: GigaChatBackgroundTaskStartRequest) -> GigaChatBackgroundTaskSummary:
        task_id = datetime.now(timezone.utc).strftime("%Y%m%d%H%M%S") + "-" + uuid.uuid4().hex[:8]
        task_dir = self.background_dir / task_id
        task_dir.mkdir(parents=True, exist_ok=True)
        async_workers = self._coerce_async_workers(req.async_workers)
        summary = GigaChatBackgroundTaskSummary(
            task_id=task_id,
            status="queued",
            filename=req.filename,
            sheet_name=req.sheet_name,
            created_at=self._now(),
            total_rows=len(req.rows),
            async_workers=async_workers,
            current_label=f"Задача поставлена в очередь: {async_workers} workers",
        )
        (task_dir / "input.json").write_text(req.model_dump_json(indent=2), encoding="utf-8")
        self._write_background_summary(summary)

        cancel_flag = threading.Event()
        worker = threading.Thread(
            target=self._run_background_labeling,
            args=(task_id, req, cancel_flag),
            daemon=True,
            name=f"gigachat-bg-{task_id}",
        )
        with self._background_lock:
            self._background_cancel_flags[task_id] = cancel_flag
            self._background_threads[task_id] = worker
        worker.start()
        return summary

    def cancel_background_task(self, task_id: str) -> GigaChatBackgroundTaskSummary:
        summary = self._read_background_summary(task_id)
        with self._background_lock:
            flag = self._background_cancel_flags.get(task_id)
            if flag:
                flag.set()
        if summary.status in {"queued", "running"}:
            summary.status = "cancelled"
            summary.finished_at = self._now()
            summary.current_label = "Отмена запрошена"
            self._write_background_summary(summary)
        return summary

    def load_background_result(self, task_id: str) -> GigaChatBackgroundTaskResultResponse:
        summary = self._read_background_summary(task_id)
        result_path = self.background_dir / task_id / "result.json"
        if not result_path.exists():
            raise FileNotFoundError(f"Background task result not found: {task_id}")
        payload = json.loads(result_path.read_text(encoding="utf-8"))
        return GigaChatBackgroundTaskResultResponse(
            task=summary,
            workbook=GigaChatWorkbookSheetDataResponse(**payload["workbook"]),
            row_runs=[GigaChatBackgroundTaskRowRun(**row) for row in payload.get("row_runs", [])],
        )

    def _run_background_labeling_row(
        self,
        req: GigaChatBackgroundTaskStartRequest,
        item: GigaChatBackgroundTaskInputRow,
        cancel_flag: threading.Event,
    ) -> dict[str, Any]:
        if cancel_flag.is_set():
            raise _BackgroundTaskCancelled()
        primary_values = (
            self._values_without_reclassification(req.values)
            if req.reclassification_enabled
            else req.values
        )
        result = self.run_row_prompt(GigaChatLabRowRunRequest(
            transport=req.transport,
            values=primary_values,
            columns=req.columns,
            row=item.source_row,
            payload_override=None if req.reclassification_enabled else req.payload_override,
            count_tokens=req.count_tokens,
        ))
        if cancel_flag.is_set():
            raise _BackgroundTaskCancelled()
        reclassification_result = None
        if req.reclassification_enabled:
            reclassification_result = self.run_row_prompt(GigaChatLabRowRunRequest(
                transport=req.transport,
                values=req.values,
                columns=req.columns,
                row=item.source_row,
                payload_override=None,
                count_tokens=req.count_tokens,
                reclassification_only=True,
            ))
        return {
            "row_index": item.row_index,
            "source_row": item.source_row,
            "result": result.model_dump(mode="json"),
            "reclassification_result": reclassification_result.model_dump(mode="json") if reclassification_result else None,
            "error": None,
        }

    def _run_background_labeling(self, task_id: str, req: GigaChatBackgroundTaskStartRequest, cancel_flag: threading.Event) -> None:
        summary = self._read_background_summary(task_id)
        task_dir = self.background_dir / task_id
        source_rows = [item.source_row for item in req.rows]
        row_runs_by_position: dict[int, dict[str, Any]] = {}
        async_workers = self._coerce_async_workers(req.async_workers)
        summary.status = "running"
        summary.started_at = self._now()
        summary.async_workers = async_workers
        summary.current_label = f"Фоновая разметка запущена: {async_workers} workers"
        self._write_background_summary(summary)

        try:
            if req.rows:
                max_workers = min(async_workers, len(req.rows))
                with concurrent.futures.ThreadPoolExecutor(max_workers=max_workers) as executor:
                    futures = {
                        executor.submit(self._run_background_labeling_row, req, item, cancel_flag): index
                        for index, item in enumerate(req.rows)
                    }
                    for future in concurrent.futures.as_completed(futures):
                        index = futures[future]
                        item = req.rows[index]
                        try:
                            row_runs_by_position[index] = future.result()
                        except _BackgroundTaskCancelled:
                            summary.status = "cancelled"
                            summary.current_label = "Задача отменена"
                            for pending in futures:
                                pending.cancel()
                            break
                        except Exception as exc:
                            summary.failed_rows += 1
                            row_runs_by_position[index] = {
                                "row_index": item.row_index,
                                "source_row": item.source_row,
                                "result": None,
                                "reclassification_result": None,
                                "error": str(exc),
                            }
                        summary.completed_rows = len(row_runs_by_position)
                        summary.progress = summary.completed_rows / summary.total_rows if summary.total_rows else 1
                        summary.current_label = (
                            f"Готово {summary.completed_rows} из {summary.total_rows}; "
                            f"workers: {max_workers}"
                        )
                        self._write_background_summary(summary)
                        if cancel_flag.is_set():
                            summary.status = "cancelled"
                            summary.current_label = "Задача отменена"
                            for pending in futures:
                                pending.cancel()
                            break
            else:
                summary.completed_rows = 0
                summary.progress = 1

            if summary.status != "cancelled":
                summary.status = "completed"
                summary.current_label = "Готово"
            summary.finished_at = self._now()
            row_runs = [row_runs_by_position[index] for index in sorted(row_runs_by_position)]
            workbook = GigaChatWorkbookSheetDataResponse(
                upload_id=f"background-{task_id}",
                filename=req.filename,
                file_format="csv",
                sheet_name=req.sheet_name,
                total_rows=len(source_rows),
                rendered_rows=len(source_rows),
                columns=req.columns,
                rows=source_rows,
            )
            result_path = task_dir / "result.json"
            result_path.write_text(json.dumps({
                "workbook": workbook.model_dump(mode="json"),
                "row_runs": row_runs,
            }, ensure_ascii=False, indent=2), encoding="utf-8")
            summary.result_path = str(result_path)
            self._write_background_summary(summary)
        except Exception as exc:
            summary.status = "failed"
            summary.error = f"{exc}\n{traceback.format_exc()}"
            summary.current_label = "Задача завершилась ошибкой"
            summary.finished_at = self._now()
            self._write_background_summary(summary)
        finally:
            with self._background_lock:
                self._background_cancel_flags.pop(task_id, None)
                self._background_threads.pop(task_id, None)

    def _read_background_summary(self, task_id: str) -> GigaChatBackgroundTaskSummary:
        meta_path = self.background_dir / task_id / "task.json"
        if not meta_path.exists():
            raise FileNotFoundError(f"Background task not found: {task_id}")
        return GigaChatBackgroundTaskSummary(**json.loads(meta_path.read_text(encoding="utf-8")))

    def _write_background_summary(self, summary: GigaChatBackgroundTaskSummary) -> None:
        task_dir = self.background_dir / summary.task_id
        task_dir.mkdir(parents=True, exist_ok=True)
        (task_dir / "task.json").write_text(summary.model_dump_json(indent=2), encoding="utf-8")

    def export_annotated_workbook(self, req: GigaChatAnnotatedExportRequest) -> tuple[str, bytes]:
        source_columns = [str(column).strip() for column in req.source_columns if str(column).strip()]
        ordered_rows = sorted(
            list(req.rows),
            key=lambda row: (row.row_index is None, row.row_index if row.row_index is not None else 0),
        )
        model_columns = [
            "Класс",
            "Новая подтематика",
            "Теги",
            "Local tags",
            "Model added tags",
            "Model rejected tags",
            "Match type",
            "Evidence",
            "Tag decisions",
            "Rule hits",
            "Suggested topics",
            "Confirmed rule hits",
            "Rejected rule hits",
            "Rule decision",
            "Model decision",
            "Reclassified topic",
            "Final topic",
            "Decision source",
        ]
        source_export_columns = self._source_export_columns(source_columns, model_columns)
        export_rows: list[dict[str, Any]] = []
        reclassified_export_rows: set[int] = set()
        for row in ordered_rows:
            new_class = self._normalize_reclassification_topic_value(row.new_class)
            is_reclassified = bool(row.is_reclassified and new_class)
            export_row: dict[str, Any] = {
                "Класс": row.classification,
                "Новая подтематика": new_class if is_reclassified else "",
                "Теги": ", ".join([str(tag).strip() for tag in row.tags if str(tag).strip()]),
                "Local tags": ", ".join([str(tag).strip() for tag in row.local_tags if str(tag).strip()]),
                "Model added tags": ", ".join([str(tag).strip() for tag in row.model_added_tags if str(tag).strip()]),
                "Model rejected tags": ", ".join([str(tag).strip() for tag in row.model_rejected_tags if str(tag).strip()]),
                "Match type": row.match_type or "",
                "Evidence": row.evidence or "",
                "Tag decisions": row.tag_decisions or "",
                "Rule hits": ", ".join([str(hit).strip() for hit in row.rule_hits if str(hit).strip()]),
                "Suggested topics": ", ".join([str(topic).strip() for topic in row.suggested_topics if str(topic).strip()]),
                "Confirmed rule hits": ", ".join([str(hit).strip() for hit in row.confirmed_rule_hits if str(hit).strip()]),
                "Rejected rule hits": ", ".join([str(hit).strip() for hit in row.rejected_rule_hits if str(hit).strip()]),
                "Rule decision": row.rule_decision or "",
                "Model decision": row.model_decision or "",
                "Reclassified topic": row.reclassified_topic or "",
                "Final topic": row.final_topic or "",
                "Decision source": row.decision_source or "",
            }
            for source_column, export_column in source_export_columns:
                export_row[export_column] = self._normalize_export_cell(row.source_row.get(source_column))
            export_rows.append(export_row)
            if is_reclassified:
                reclassified_export_rows.add(len(export_rows) + 1)

        ordered_columns = [
            *model_columns,
            *[export_column for _, export_column in source_export_columns],
        ]
        df = pd.DataFrame(export_rows, columns=ordered_columns)
        df = self._sanitize_for_excel(df)

        buffer = BytesIO()
        sheet_name = self._excel_safe_sheet_name(req.sheet_name or "Разметка")
        with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
            df.to_excel(writer, sheet_name=sheet_name, index=False)
            ws = writer.sheets[sheet_name]
            ws.freeze_panes = "A2"
            ws.auto_filter.ref = ws.dimensions
            reclassified_fill = PatternFill(fill_type="solid", fgColor="FEF3C7")
            for export_row_idx in reclassified_export_rows:
                ws.cell(row=export_row_idx, column=2).fill = reclassified_fill
            for col_idx in range(1, ws.max_column + 1):
                for row_idx in range(2, ws.max_row + 1):
                    cell = ws.cell(row=row_idx, column=col_idx)
                    if isinstance(cell.value, datetime):
                        cell.number_format = "yyyy-mm-dd hh:mm:ss"
                    elif isinstance(cell.value, date):
                        cell.number_format = "yyyy-mm-dd"

        buffer.seek(0)
        export_filename = f"{Path(req.filename or 'annotated.xlsx').stem}_annotated.xlsx"
        return export_filename, buffer.getvalue()

    def export_validation_workbook(self, req: GigaChatAnnotatedExportRequest) -> tuple[str, bytes]:
        source_columns = [str(column).strip() for column in req.source_columns if str(column).strip()]
        ordered_rows = sorted(
            list(req.rows),
            key=lambda row: (row.row_index is None, row.row_index if row.row_index is not None else 0),
        )
        model_columns = [
            "Review status",
            "Class matches",
            "Tags match",
            "Topic matches",
            "Expected class",
            "Expected tags",
            "Expected final topic",
            "Reviewer comment",
            "Model class",
            "Model tags",
            "Local tags",
            "Model added tags",
            "Model rejected tags",
            "Match type",
            "Evidence",
            "Tag decisions",
            "Rule hits",
            "Suggested topics",
            "Confirmed rule hits",
            "Rejected rule hits",
            "Rule decision",
            "Model decision",
            "Reclassified topic",
            "Model final topic",
            "Decision source",
        ]
        source_export_columns = self._source_export_columns(source_columns, model_columns)
        export_rows: list[dict[str, Any]] = []
        for row in ordered_rows:
            model_class = row.classification or ""
            model_tags = ", ".join([str(tag).strip() for tag in row.tags if str(tag).strip()])
            model_topic = row.final_topic or row.reclassified_topic or row.classification or ""
            export_row: dict[str, Any] = {
                "Review status": "needs_review",
                "Class matches": "",
                "Tags match": "",
                "Topic matches": "",
                "Expected class": "",
                "Expected tags": "",
                "Expected final topic": "",
                "Reviewer comment": "",
                "Model class": model_class,
                "Model tags": model_tags,
                "Local tags": ", ".join([str(tag).strip() for tag in row.local_tags if str(tag).strip()]),
                "Model added tags": ", ".join([str(tag).strip() for tag in row.model_added_tags if str(tag).strip()]),
                "Model rejected tags": ", ".join([str(tag).strip() for tag in row.model_rejected_tags if str(tag).strip()]),
                "Match type": row.match_type or "",
                "Evidence": row.evidence or "",
                "Tag decisions": row.tag_decisions or "",
                "Rule hits": ", ".join([str(hit).strip() for hit in row.rule_hits if str(hit).strip()]),
                "Suggested topics": ", ".join([str(topic).strip() for topic in row.suggested_topics if str(topic).strip()]),
                "Confirmed rule hits": ", ".join([str(hit).strip() for hit in row.confirmed_rule_hits if str(hit).strip()]),
                "Rejected rule hits": ", ".join([str(hit).strip() for hit in row.rejected_rule_hits if str(hit).strip()]),
                "Rule decision": row.rule_decision or "",
                "Model decision": row.model_decision or "",
                "Reclassified topic": row.reclassified_topic or "",
                "Model final topic": model_topic,
                "Decision source": row.decision_source or "",
            }
            for source_column, export_column in source_export_columns:
                export_row[export_column] = self._normalize_export_cell(row.source_row.get(source_column))
            export_rows.append(export_row)

        ordered_columns = [
            *model_columns,
            *[export_column for _, export_column in source_export_columns],
        ]
        df = pd.DataFrame(export_rows, columns=ordered_columns)
        df = self._sanitize_for_excel(df)

        buffer = BytesIO()
        sheet_name = self._excel_safe_sheet_name(req.sheet_name or "Validation")
        with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
            df.to_excel(writer, sheet_name=sheet_name, index=False)
            ws = writer.sheets[sheet_name]
            ws.freeze_panes = "I2"
            ws.auto_filter.ref = ws.dimensions

            from openpyxl.styles import Font, PatternFill
            from openpyxl.worksheet.datavalidation import DataValidation

            manual_fill = PatternFill("solid", fgColor="FFF2CC")
            model_fill = PatternFill("solid", fgColor="D9EAF7")
            source_fill = PatternFill("solid", fgColor="E2F0D9")
            for col_idx in range(1, ws.max_column + 1):
                header = ws.cell(row=1, column=col_idx)
                header.font = Font(bold=True)
                if col_idx <= 8:
                    header.fill = manual_fill
                elif col_idx <= 19:
                    header.fill = model_fill
                else:
                    header.fill = source_fill
                for row_idx in range(2, ws.max_row + 1):
                    cell = ws.cell(row=row_idx, column=col_idx)
                    if isinstance(cell.value, datetime):
                        cell.number_format = "yyyy-mm-dd hh:mm:ss"
                    elif isinstance(cell.value, date):
                        cell.number_format = "yyyy-mm-dd"

            if ws.max_row >= 2:
                yes_no = DataValidation(type="list", formula1='"yes,no,needs_review"', allow_blank=True)
                ws.add_data_validation(yes_no)
                for col_letter in ("B", "C", "D"):
                    yes_no.add(f"{col_letter}2:{col_letter}{ws.max_row}")
                status = DataValidation(type="list", formula1='"needs_review,approved,rejected"', allow_blank=True)
                ws.add_data_validation(status)
                status.add(f"A2:A{ws.max_row}")

            for col_idx in range(1, ws.max_column + 1):
                max_len = len(str(ws.cell(row=1, column=col_idx).value or ""))
                for row_idx in range(2, min(ws.max_row, 40) + 1):
                    max_len = max(max_len, len(str(ws.cell(row=row_idx, column=col_idx).value or "")))
                ws.column_dimensions[ws.cell(row=1, column=col_idx).column_letter].width = max(12, min(48, max_len * 0.8))

        buffer.seek(0)
        export_filename = f"{Path(req.filename or 'annotated.xlsx').stem}_validation.xlsx"
        return export_filename, buffer.getvalue()

    def export_workbook_rows(self, req: GigaChatWorkbookRowsExportRequest) -> tuple[str, bytes]:
        source_columns = [str(column).strip() for column in req.source_columns if str(column).strip()]
        ordered_rows = sorted(
            list(req.rows),
            key=lambda row: (row.row_index is None, row.row_index if row.row_index is not None else 0),
        )
        model_columns = [
            "Row index",
            "Rule hits",
            "Suggested action",
            "Matched keywords",
            "Matched fields",
            "Suggested topic",
        ]
        source_export_columns = self._source_export_columns(source_columns, model_columns)
        export_rows: list[dict[str, Any]] = []
        for row in ordered_rows:
            export_row: dict[str, Any] = {
                "Row index": "" if row.row_index is None else int(row.row_index) + 1,
                "Rule hits": ", ".join([str(hit).strip() for hit in row.rule_hits if str(hit).strip()]),
                "Suggested action": ", ".join([str(item).strip() for item in row.suggested_actions if str(item).strip()]),
                "Matched keywords": ", ".join([str(item).strip() for item in row.matched_keywords if str(item).strip()]),
                "Matched fields": ", ".join([str(item).strip() for item in row.matched_fields if str(item).strip()]),
                "Suggested topic": ", ".join([str(item).strip() for item in row.suggested_topics if str(item).strip()]),
            }
            for source_column, export_column in source_export_columns:
                export_row[export_column] = self._normalize_export_cell(row.source_row.get(source_column))
            export_rows.append(export_row)

        ordered_columns = [
            *model_columns,
            *[export_column for _, export_column in source_export_columns],
        ]
        df = pd.DataFrame(export_rows, columns=ordered_columns)
        df = self._sanitize_for_excel(df)

        buffer = BytesIO()
        sheet_name = self._excel_safe_sheet_name(req.sheet_name or "data")
        with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
            df.to_excel(writer, sheet_name=sheet_name, index=False)
            ws = writer.sheets[sheet_name]
            ws.freeze_panes = "A2"
            ws.auto_filter.ref = ws.dimensions
            for col_idx in range(1, ws.max_column + 1):
                for row_idx in range(2, ws.max_row + 1):
                    cell = ws.cell(row=row_idx, column=col_idx)
                    if isinstance(cell.value, datetime):
                        cell.number_format = "yyyy-mm-dd hh:mm:ss"
                    elif isinstance(cell.value, date):
                        cell.number_format = "yyyy-mm-dd"

        buffer.seek(0)
        export_filename = f"{Path(req.filename or 'workbook.xlsx').stem}_{req.sheet_name or 'sheet'}_rows.xlsx"
        return export_filename, buffer.getvalue()

    def start_chunked_workbook_upload(
        self,
        req: GigaChatWorkbookChunkedUploadStartRequest,
        *,
        user_id: str | None = None,
        workspace_id: str | None = None,
    ) -> GigaChatWorkbookChunkedUploadStartResponse:
        if self.catalog and (not user_id or not workspace_id):
            user = self.catalog.get_or_create_dev_user()
            user_id = user_id or str(user.get("id") or "dev")
            workspace_id = workspace_id or str(user.get("workspace_id") or "default")
        filename = Path(req.filename or "upload.xlsx").name or "upload.xlsx"
        total_size = max(0, int(req.total_size))
        chunk_size = max(1024 * 1024, min(int(req.chunk_size or 0), 64 * 1024 * 1024))
        total_chunks = max(1, int(req.total_chunks or math.ceil(total_size / chunk_size) or 1))
        session_id = datetime.now(timezone.utc).strftime("%Y%m%d%H%M%S") + "-" + uuid.uuid4().hex[:8]
        session_dir = self.chunked_uploads_dir / session_id
        (session_dir / "chunks").mkdir(parents=True, exist_ok=True)
        payload = {
            "session_id": session_id,
            "filename": filename,
            "total_size": total_size,
            "chunk_size": chunk_size,
            "total_chunks": total_chunks,
            "received_chunks": [],
            "received_bytes": 0,
            "user_id": user_id or "anonymous",
            "workspace_id": workspace_id or "default",
            "created_at": self._now().isoformat(),
        }
        self._write_chunk_session(payload)
        return GigaChatWorkbookChunkedUploadStartResponse(
            session_id=session_id,
            filename=filename,
            chunk_size=chunk_size,
            total_chunks=total_chunks,
        )

    def receive_chunked_workbook_chunk(
        self,
        session_id: str,
        chunk_index: int,
        content: bytes,
    ) -> GigaChatWorkbookChunkUploadResponse:
        with self._upload_task_lock:
            session = self._read_chunk_session(session_id)
            if session.get("status") == "cancelled":
                raise ValueError("Upload session was cancelled.")
            total_chunks = int(session["total_chunks"])
            if chunk_index < 0 or chunk_index >= total_chunks:
                raise ValueError(f"Chunk index is out of range: {chunk_index}")
            chunk_dir = self.chunked_uploads_dir / session_id / "chunks"
            chunk_dir.mkdir(parents=True, exist_ok=True)
            chunk_path = chunk_dir / f"{chunk_index:08d}.part"
            already_received = chunk_path.exists()
            chunk_path.write_bytes(content)
            received = set(int(item) for item in session.get("received_chunks", []))
            received.add(chunk_index)
            session["received_chunks"] = sorted(received)
            if already_received:
                session["received_bytes"] = sum((chunk_dir / f"{idx:08d}.part").stat().st_size for idx in received)
            else:
                session["received_bytes"] = int(session.get("received_bytes") or 0) + len(content)
            self._write_chunk_session(session)
        return GigaChatWorkbookChunkUploadResponse(
            session_id=session_id,
            chunk_index=chunk_index,
            total_chunks=total_chunks,
            received_chunks=len(session["received_chunks"]),
            received_bytes=int(session["received_bytes"]),
            total_size=int(session.get("total_size") or 0),
        )

    def complete_chunked_workbook_upload(self, session_id: str) -> GigaChatWorkbookChunkedUploadCompleteResponse:
        session = self._read_chunk_session(session_id)
        if session.get("status") == "cancelled":
            raise ValueError("Upload session was cancelled.")
        received = set(int(item) for item in session.get("received_chunks", []))
        total_chunks = int(session["total_chunks"])
        missing = [idx for idx in range(total_chunks) if idx not in received]
        if missing:
            raise ValueError(f"Upload is incomplete. Missing chunks: {missing[:10]}")
        task_id = datetime.now(timezone.utc).strftime("%Y%m%d%H%M%S") + "-" + uuid.uuid4().hex[:8]
        task = {
            "task_id": task_id,
            "session_id": session_id,
            "status": "queued",
            "phase": "queued",
            "message": "Файл загружен, задача поставлена в очередь",
            "progress": 0.45,
            "total_size": int(session.get("total_size") or 0),
            "received_bytes": int(session.get("received_bytes") or 0),
            "total_chunks": total_chunks,
            "received_chunks": len(received),
            "error": None,
            "workbook": None,
            "created_at": self._now().isoformat(),
            "updated_at": self._now().isoformat(),
        }
        self._write_upload_task(task)
        cancel_flag = threading.Event()
        worker = threading.Thread(target=self._run_chunked_upload_task, args=(task_id, cancel_flag), daemon=True)
        with self._upload_task_lock:
            self._upload_cancel_flags[task_id] = cancel_flag
            self._upload_task_threads[task_id] = worker
        worker.start()
        return GigaChatWorkbookChunkedUploadCompleteResponse(task_id=task_id, session_id=session_id, status="queued")

    def cancel_chunked_workbook_upload_session(self, session_id: str) -> dict[str, Any]:
        with self._upload_task_lock:
            session = self._read_chunk_session(session_id)
            session["status"] = "cancelled"
            session["cancelled_at"] = self._now().isoformat()
            self._write_chunk_session(session)
        with suppress(Exception):
            self._cleanup_chunk_session_files(session_id)
        return {"session_id": session_id, "status": "cancelled"}

    def get_workbook_upload_task(self, task_id: str) -> GigaChatWorkbookUploadTaskResponse:
        task = self._read_upload_task(task_id)
        workbook = task.get("workbook")
        return GigaChatWorkbookUploadTaskResponse(
            task_id=str(task["task_id"]),
            session_id=str(task["session_id"]),
            status=task["status"],
            phase=str(task.get("phase") or ""),
            message=str(task.get("message") or ""),
            progress=float(task.get("progress") or 0),
            total_size=int(task.get("total_size") or 0),
            received_bytes=int(task.get("received_bytes") or 0),
            total_chunks=int(task.get("total_chunks") or 0),
            received_chunks=int(task.get("received_chunks") or 0),
            error=task.get("error"),
            workbook=GigaChatWorkbookUploadResponse(**workbook) if isinstance(workbook, dict) else None,
        )

    def cancel_workbook_upload_task(self, task_id: str) -> GigaChatWorkbookUploadTaskResponse:
        with self._upload_task_lock:
            task = self._read_upload_task(task_id)
            if task.get("status") not in {"completed", "failed", "cancelled"}:
                flag = self._upload_cancel_flags.get(task_id)
                if flag:
                    flag.set()
                task.update({
                    "status": "cancelled",
                    "phase": "cancelled",
                    "message": "Загрузка отменена пользователем",
                    "error": None,
                })
                self._write_upload_task(task)
                with suppress(Exception):
                    self._cleanup_chunk_session_files(str(task.get("session_id") or ""))
        return self.get_workbook_upload_task(task_id)

    def _run_chunked_upload_task(self, task_id: str, cancel_flag: threading.Event) -> None:
        task = self._read_upload_task(task_id)
        session = self._read_chunk_session(str(task["session_id"]))
        session_id = str(session["session_id"])
        upload_id = datetime.now(timezone.utc).strftime("%Y%m%d%H%M%S") + "-" + uuid.uuid4().hex[:8]
        upload_dir = self.uploads_dir / upload_id
        upload_dir.mkdir(parents=True, exist_ok=True)
        suffix = (Path(str(session["filename"])).suffix or ".xlsx").lower()
        assembled_path = upload_dir / f"source{suffix}"
        try:
            self._raise_if_upload_cancelled(task_id, cancel_flag)
            self._update_upload_task(task_id, status="running", phase="assembling", progress=0.50, message="Собираем файл из загруженных частей")
            chunk_dir = self.chunked_uploads_dir / session_id / "chunks"
            with assembled_path.open("wb") as target:
                for idx in range(int(session["total_chunks"])):
                    self._raise_if_upload_cancelled(task_id, cancel_flag)
                    chunk_path = chunk_dir / f"{idx:08d}.part"
                    if not chunk_path.exists():
                        raise FileNotFoundError(f"Missing upload chunk: {idx}")
                    with chunk_path.open("rb") as source:
                        while True:
                            self._raise_if_upload_cancelled(task_id, cancel_flag)
                            block = source.read(1024 * 1024)
                            if not block:
                                break
                            target.write(block)
            self._raise_if_upload_cancelled(task_id, cancel_flag)
            self._update_upload_task(task_id, phase="parsing", progress=0.62, message="Читаем структуру Excel/CSV и готовим превью")
            workbook = self._finalize_workbook_upload(
                upload_id=upload_id,
                original_filename=str(session["filename"]),
                display_filename=str(session["filename"]),
                stored_path=assembled_path,
                user_id=str(session.get("user_id") or "anonymous"),
                workspace_id=str(session.get("workspace_id") or "default"),
                progress_callback=lambda phase, progress, message: self._update_upload_task_checked(task_id, cancel_flag, phase=phase, progress=progress, message=message),
            )
            self._raise_if_upload_cancelled(task_id, cancel_flag)
            self._update_upload_task(
                task_id,
                status="completed",
                phase="completed",
                progress=1,
                message="Файл готов",
                workbook=workbook.model_dump(mode="json"),
            )
        except _WorkbookUploadCancelled:
            with suppress(Exception):
                if self.lake:
                    self.lake.delete_file_artifacts(
                        file_id=upload_id,
                        workspace_id=str(session.get("workspace_id") or "default"),
                    )
                elif self.catalog:
                    self.catalog.delete_file(upload_id)
            self._update_upload_task(
                task_id,
                status="cancelled",
                phase="cancelled",
                message="Загрузка отменена пользователем",
                error=None,
            )
            with suppress(Exception):
                shutil.rmtree(upload_dir)
        except Exception as exc:
            self._update_upload_task(
                task_id,
                status="failed",
                phase="failed",
                progress=float(task.get("progress") or 0),
                message="Не удалось обработать файл",
                error=f"{exc}\n{traceback.format_exc()}",
            )
        finally:
            with suppress(Exception):
                self._cleanup_chunk_session_files(session_id)
            with self._upload_task_lock:
                self._upload_cancel_flags.pop(task_id, None)
                self._upload_task_threads.pop(task_id, None)

    def _chunk_session_path(self, session_id: str) -> Path:
        return self.chunked_uploads_dir / session_id / "session.json"

    def _upload_task_path(self, task_id: str) -> Path:
        return self.chunked_uploads_dir / "tasks" / f"{task_id}.json"

    def _cleanup_chunk_session_files(self, session_id: str) -> None:
        if session_id:
            shutil.rmtree(self.chunked_uploads_dir / session_id / "chunks")

    def _read_chunk_session(self, session_id: str) -> dict[str, Any]:
        path = self._chunk_session_path(session_id)
        if not path.exists():
            raise FileNotFoundError(f"Chunked upload session not found: {session_id}")
        return json.loads(path.read_text(encoding="utf-8"))

    def _write_chunk_session(self, payload: dict[str, Any]) -> None:
        session_id = str(payload["session_id"])
        path = self._chunk_session_path(session_id)
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")

    def _read_upload_task(self, task_id: str) -> dict[str, Any]:
        path = self._upload_task_path(task_id)
        if not path.exists():
            raise FileNotFoundError(f"Workbook upload task not found: {task_id}")
        return json.loads(path.read_text(encoding="utf-8"))

    def _write_upload_task(self, payload: dict[str, Any]) -> None:
        path = self._upload_task_path(str(payload["task_id"]))
        path.parent.mkdir(parents=True, exist_ok=True)
        payload["updated_at"] = self._now().isoformat()
        path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")

    def _update_upload_task(self, task_id: str, **updates: Any) -> None:
        task = self._read_upload_task(task_id)
        if task.get("status") == "cancelled" and updates.get("status") not in {"cancelled", None}:
            return
        task.update(updates)
        self._write_upload_task(task)

    def _update_upload_task_checked(self, task_id: str, cancel_flag: threading.Event, **updates: Any) -> None:
        self._raise_if_upload_cancelled(task_id, cancel_flag)
        self._update_upload_task(task_id, **updates)
        self._raise_if_upload_cancelled(task_id, cancel_flag)

    def _raise_if_upload_cancelled(self, task_id: str, cancel_flag: threading.Event) -> None:
        if cancel_flag.is_set():
            raise _WorkbookUploadCancelled()
        try:
            task = self._read_upload_task(task_id)
        except Exception:
            return
        if task.get("status") == "cancelled":
            cancel_flag.set()
            raise _WorkbookUploadCancelled()

    def _finalize_workbook_upload(
        self,
        *,
        upload_id: str,
        original_filename: str,
        display_filename: str,
        stored_path: Path,
        user_id: str,
        workspace_id: str,
        progress_callback: Any | None = None,
    ) -> GigaChatWorkbookUploadResponse:
        if stored_path.suffix.lower() == ".zip":
            archive_path = stored_path.with_name("archive.zip")
            stored_path.rename(archive_path)
            inner_name, source_content = self._extract_supported_file_from_zip(original_filename, archive_path.read_bytes())
            inner_ext = (Path(inner_name).suffix or ".xlsx").lower()
            stored_path = archive_path.with_name(f"source{inner_ext}")
            stored_path.write_bytes(source_content)
            display_filename = f"{original_filename} -> {inner_name}"

        progress_callback and progress_callback("parsing", 0.64, "Читаем листы и первые строки")
        meta = self._inspect_workbook(stored_path, display_filename)
        if self.catalog:
            self.catalog.create_file(
                file_id=upload_id,
                workspace_id=workspace_id,
                uploaded_by_user_id=user_id,
                original_filename=original_filename,
                display_filename=display_filename,
                storage_path=str(stored_path),
                file_format=str(meta.get("file_format", "excel")),
                sheet_count=int(meta.get("sheet_count", 0)),
                metadata=meta,
            )
        if self.lake:
            try:
                progress_callback and progress_callback("ingesting", 0.78, "Пишем parquet-слой и проверяем дубли")
                lake_result = self.lake.ingest_workbook(
                    file_id=upload_id,
                    workspace_id=workspace_id,
                    user_id=user_id,
                    stored_path=stored_path,
                    display_filename=display_filename,
                    file_format=str(meta.get("file_format", "excel")),
                    sheets=list(meta.get("sheets", [])),
                )
                input_rows_by_sheet = {
                    str(artifact.get("metadata", {}).get("sheet_name") or ""): int(artifact.get("metadata", {}).get("input_rows") or 0)
                    for artifact in lake_result.get("artifacts", [])
                    if isinstance(artifact, dict)
                }
                for sheet in meta.get("sheets", []):
                    if not isinstance(sheet, dict):
                        continue
                    input_rows = input_rows_by_sheet.get(str(sheet.get("name") or ""))
                    if input_rows is not None:
                        sheet["rows_total"] = input_rows
            except Exception as exc:
                if self.catalog:
                    self.catalog.update_file_status(upload_id, "ingest_failed")
                meta["lake_error"] = str(exc)
        (stored_path.parent / "meta.json").write_text(json.dumps(meta, ensure_ascii=False, indent=2), encoding="utf-8")
        return GigaChatWorkbookUploadResponse(**meta)

    def upload_workbook(
        self,
        filename: str,
        content: bytes,
        *,
        user_id: str | None = None,
        workspace_id: str | None = None,
    ) -> GigaChatWorkbookUploadResponse:
        upload_id = datetime.now(timezone.utc).strftime("%Y%m%d%H%M%S") + "-" + uuid.uuid4().hex[:8]
        upload_dir = self.uploads_dir / upload_id
        upload_dir.mkdir(parents=True, exist_ok=True)
        if self.catalog and (not user_id or not workspace_id):
            user = self.catalog.get_or_create_dev_user()
            user_id = user_id or str(user.get("id") or "dev")
            workspace_id = workspace_id or str(user.get("workspace_id") or "default")
        user_id = user_id or "anonymous"
        workspace_id = workspace_id or "default"
        ext = (Path(filename).suffix or ".xlsx").lower()
        stored_path = upload_dir / f"source{ext}"
        stored_path.write_bytes(content)
        return self._finalize_workbook_upload(
            upload_id=upload_id,
            original_filename=filename,
            display_filename=filename,
            stored_path=stored_path,
            user_id=user_id,
            workspace_id=workspace_id,
        )

    def upload_local_workbook(self, filename: str, *, user_id: str | None = None, workspace_id: str | None = None) -> GigaChatWorkbookUploadResponse:
        requested_name = Path(filename).name
        if not requested_name or requested_name != filename:
            raise ValueError("Local workbook fallback accepts only a file name, not a path.")
        allowed_dirs = [
            Path("data/raw"),
            Path("outputs"),
            Path("outputs/gigachat_delivery"),
        ]
        for directory in allowed_dirs:
            candidate = directory / requested_name
            if candidate.is_file():
                return self.upload_workbook(candidate.name, candidate.read_bytes(), user_id=user_id, workspace_id=workspace_id)
        searched = ", ".join(str(directory / requested_name) for directory in allowed_dirs)
        raise FileNotFoundError(f"Local workbook not found. Checked: {searched}")

    def load_sheet(self, upload_id: str, req: GigaChatWorkbookSelectSheetRequest) -> GigaChatWorkbookSheetDataResponse:
        upload_dir = self.uploads_dir / upload_id
        meta_path = upload_dir / "meta.json"
        if not meta_path.exists():
            raise FileNotFoundError(f"Workbook upload not found: {upload_id}")
        meta = json.loads(meta_path.read_text(encoding="utf-8"))
        filename = str(meta.get("filename", ""))
        file_format = str(meta.get("file_format", "excel"))
        stored_path = next(upload_dir.glob("source.*"), None)
        if stored_path is None:
            raise FileNotFoundError(f"Workbook source file not found for upload: {upload_id}")

        meta_sheets = [sheet for sheet in meta.get("sheets", []) if isinstance(sheet, dict)]
        meta_sheet = next((sheet for sheet in meta_sheets if str(sheet.get("name", "")) == (req.sheet_name if file_format != "csv" else "data")), None)
        source_total_rows = int(meta_sheet.get("rows_total") or 0) if meta_sheet else 0
        source_preview_rows = len(meta_sheet.get("preview_rows") or []) if meta_sheet else 0

        if self.lake:
            lake_sheet = self.lake.load_raw_sheet(upload_id, req.sheet_name if file_format != "csv" else "data", req.row_limit)
            if lake_sheet is not None:
                lake_total_rows = int(lake_sheet.get("total_rows") or 0)
                lake_has_complete_source_rows = lake_total_rows > 0 and source_total_rows > 0 and lake_total_rows >= source_total_rows
                if lake_has_complete_source_rows:
                    return GigaChatWorkbookSheetDataResponse(
                        upload_id=upload_id,
                        filename=filename,
                        file_format="csv" if file_format == "csv" else "excel",
                        sheet_name=req.sheet_name if file_format != "csv" else "data",
                        total_rows=lake_total_rows,
                        rendered_rows=int(lake_sheet["rendered_rows"]),
                        columns=list(lake_sheet["columns"]),
                        rows=list(lake_sheet["rows"]),
                    )
                if lake_total_rows <= 0 and source_total_rows <= 0 and source_preview_rows <= 0:
                    return GigaChatWorkbookSheetDataResponse(
                        upload_id=upload_id,
                        filename=filename,
                        file_format="csv" if file_format == "csv" else "excel",
                        sheet_name=req.sheet_name if file_format != "csv" else "data",
                        total_rows=0,
                        rendered_rows=0,
                        columns=list(lake_sheet["columns"]),
                        rows=[],
                    )
        if file_format == "csv":
            df = self._read_csv_frame(stored_path)
            sheet_name = "data"
        else:
            sheet_name = req.sheet_name
            available_sheets = {str(sheet.get("name", "")) for sheet in meta.get("sheets", [])}
            if sheet_name not in available_sheets:
                raise ValueError(f"Sheet not found in workbook: {sheet_name}")
            df = pd.read_excel(stored_path, sheet_name=sheet_name, dtype=object)
        df = self._drop_empty_unnamed_columns(df)

        columns = [str(c) for c in df.columns]
        row_limit = min(max(int(req.row_limit), 1), len(df))
        rows = self._frame_to_rows(df.head(row_limit))
        return GigaChatWorkbookSheetDataResponse(
            upload_id=upload_id,
            filename=filename,
            file_format="csv" if file_format == "csv" else "excel",
            sheet_name=sheet_name,
            total_rows=int(len(df)),
            rendered_rows=int(len(rows)),
            columns=columns,
            rows=rows,
        )

    def _inspect_workbook(self, stored_path: Path, filename: str) -> dict[str, Any]:
        suffix = stored_path.suffix.lower()
        if suffix == ".csv":
            df = self._read_csv_frame(stored_path)
            df = self._drop_empty_unnamed_columns(df)
            sheet = self._build_sheet_preview("data", df)
            return {
                "upload_id": stored_path.parent.name,
                "filename": filename,
                "file_format": "csv",
                "sheet_count": 1,
                "sheets": [sheet.model_dump()],
            }

        if suffix == ".xls":
            excel = pd.ExcelFile(stored_path)
            sheets = [
                self._build_sheet_preview(
                    sheet_name,
                    self._drop_empty_unnamed_columns(pd.read_excel(excel, sheet_name=sheet_name, dtype=object)),
                ).model_dump()
                for sheet_name in excel.sheet_names
            ]
        else:
            sheets = self._inspect_excel_workbook_fast(stored_path)
        return {
            "upload_id": stored_path.parent.name,
            "filename": filename,
            "file_format": "excel",
            "sheet_count": len(sheets),
            "sheets": sheets,
        }

    def _inspect_excel_workbook_fast(self, stored_path: Path) -> list[dict[str, Any]]:
        from openpyxl import load_workbook

        workbook = load_workbook(stored_path, read_only=True, data_only=True)
        sheets: list[dict[str, Any]] = []
        try:
            for worksheet in workbook.worksheets:
                rows_iter = worksheet.iter_rows(values_only=True)
                try:
                    raw_header = next(rows_iter)
                except StopIteration:
                    raw_header = []
                columns = self._normalize_excel_header(raw_header)
                preview_rows: list[dict[str, Any]] = []
                for raw_row in rows_iter:
                    if len(preview_rows) >= 5:
                        break
                    values = list(raw_row[:len(columns)])
                    if len(values) < len(columns):
                        values.extend([None] * (len(columns) - len(values)))
                    preview_rows.append({column: self._normalize_cell(value) for column, value in zip(columns, values)})
                row_total = int(worksheet.max_row - 1) if worksheet.max_row else 0
                sheets.append(GigaChatWorkbookSheetPreview(
                    name=str(worksheet.title),
                    rows_total=max(0, row_total),
                    column_count=len(columns),
                    columns=columns,
                    preview_rows=preview_rows,
                ).model_dump())
        finally:
            workbook.close()
        return sheets

    @staticmethod
    def _normalize_excel_header(raw_header: Any) -> list[str]:
        seen: dict[str, int] = {}
        columns: list[str] = []
        for idx, value in enumerate(list(raw_header or []), start=1):
            name = str(value).strip() if value is not None else ""
            if not name:
                name = f"Column {idx}"
            count = seen.get(name, 0)
            seen[name] = count + 1
            columns.append(name if count == 0 else f"{name}_{count + 1}")
        return columns

    @classmethod
    def _extract_supported_file_from_zip(cls, filename: str, content: bytes) -> tuple[str, bytes]:
        supported_suffixes = (".xlsx", ".xls", ".xlsm", ".csv")
        suffix_priority = {".xlsx": 0, ".xlsm": 1, ".xls": 2, ".csv": 3}
        with ZipFile(BytesIO(content)) as archive:
            candidates = [
                info for info in archive.infolist()
                if not info.is_dir()
                and not Path(info.filename).name.startswith("._")
                and "__MACOSX/" not in info.filename
                and info.filename.lower().endswith(supported_suffixes)
            ]
            if not candidates:
                raise ValueError(
                    f"В архиве {filename} не найдено поддерживаемых файлов Excel/CSV. "
                    "Ожидается .xlsx, .xls, .xlsm или .csv."
                )
            candidates.sort(key=lambda info: (
                suffix_priority.get(Path(info.filename).suffix.lower(), 99),
                len(Path(info.filename).parts),
                info.filename.lower(),
            ))
            selected = candidates[0]
            return selected.filename, archive.read(selected)

    @classmethod
    def _read_csv_frame(cls, stored_path: Path) -> pd.DataFrame:
        errors: list[str] = []
        for encoding in ("utf-8-sig", "utf-8", "cp1251"):
            with suppress(Exception):
                return cls._normalize_csv_columns(pd.read_csv(stored_path, dtype=object, sep=None, engine="python", encoding=encoding))
            for separator in (";", ",", "\t", "|"):
                try:
                    df = pd.read_csv(stored_path, dtype=object, sep=separator, encoding=encoding)
                    if len(df.columns) > 1:
                        return cls._normalize_csv_columns(df)
                except Exception as exc:
                    errors.append(f"{encoding}/{separator}: {exc}")
        message = "; ".join(errors[:4]) or "не удалось определить кодировку или разделитель"
        raise ValueError(
            "Не удалось прочитать CSV. Проверьте разделитель и кавычки в файле. "
            f"Последние ошибки: {message}"
        )

    @staticmethod
    def _normalize_csv_columns(df: pd.DataFrame) -> pd.DataFrame:
        df = df.copy()
        df.columns = [str(column).strip() for column in df.columns]
        return df

    @staticmethod
    def _drop_empty_unnamed_columns(df: pd.DataFrame) -> pd.DataFrame:
        if df.empty and not len(df.columns):
            return df
        keep_columns = []
        for column in df.columns:
            column_name = str(column)
            is_unnamed = column_name.startswith("Unnamed:")
            if not is_unnamed:
                keep_columns.append(column)
                continue
            series = df[column]
            has_value = bool(series.map(lambda value: not pd.isna(value) and str(value).strip() != "").any())
            if has_value:
                keep_columns.append(column)
        return df.loc[:, keep_columns]

    @staticmethod
    def _split_rule_list(raw: Any) -> list[str]:
        if isinstance(raw, list):
            return [str(item).strip() for item in raw if str(item).strip()]
        if raw is None:
            return []
        text = str(raw).replace("\r", "\n")
        parts = [part.strip() for chunk in text.split("\n") for part in chunk.split(",")]
        return [part for part in parts if part]

    @staticmethod
    def _split_keyword_list(raw: Any) -> list[str]:
        if isinstance(raw, list):
            return [str(item) for item in raw if str(item).strip()]
        if raw is None:
            return []
        text = str(raw).replace("\r", "\n")
        if "," in text:
            parts = [part.strip() for chunk in text.split("\n") for part in chunk.split(",")]
            return [part for part in parts if part]
        if "\n" in text:
            return [part for part in text.split("\n") if part.strip()]
        return [text] if text.strip() else []

    @staticmethod
    def _normalize_rule_pack_import_header(raw: Any) -> str:
        text = str(raw or "").strip().lower().replace("ё", "е")
        return re.sub(r"[^0-9a-zа-я]+", "", text)

    @classmethod
    def _cell_text(cls, value: Any) -> str:
        normalized = cls._normalize_cell(value)
        if normalized is None:
            return ""
        if isinstance(normalized, float) and normalized.is_integer():
            return str(int(normalized))
        return str(normalized).strip()

    @classmethod
    def _parse_rule_pack_import_bool(cls, value: Any, *, default: bool = True) -> bool:
        text = cls._cell_text(value).strip().lower()
        if not text:
            return default
        if text in {"1", "true", "yes", "y", "on", "да", "активно", "активен"}:
            return True
        if text in {"0", "false", "no", "n", "off", "нет", "неактивно", "не активен"}:
            return False
        return default

    @classmethod
    def _parse_rule_pack_import_filters(cls, value: Any) -> list[GigaChatRulePackFilter]:
        text = cls._cell_text(value)
        if not text:
            return []
        try:
            parsed = json.loads(text)
        except Exception:
            return []
        if isinstance(parsed, dict):
            parsed = [parsed]
        if not isinstance(parsed, list):
            return []
        filters: list[GigaChatRulePackFilter] = []
        for item in parsed:
            if not isinstance(item, dict):
                continue
            field = cls._cell_text(item.get("field"))
            op = cls._cell_text(item.get("op") or "eq").lower()
            filter_value = cls._cell_text(item.get("value"))
            if not field or not filter_value or op not in {"eq", "ne"}:
                continue
            filters.append(GigaChatRulePackFilter(field=field, op=op, value=filter_value))
        return filters

    @classmethod
    def _parse_rule_pack_items(cls, raw: Any) -> list[GigaChatRulePack]:
        text = str(raw or "").strip()
        if not text:
            return []
        try:
            parsed = json.loads(text)
        except Exception:
            return []
        if not isinstance(parsed, list):
            return []

        items: list[GigaChatRulePack] = []
        for index, entry in enumerate(parsed):
            if not isinstance(entry, dict):
                continue
            filters: list[GigaChatRulePackFilter] = []
            for filter_entry in entry.get("filters", []) or []:
                if not isinstance(filter_entry, dict):
                    continue
                field = str(filter_entry.get("field", "") or "").strip()
                op = str(filter_entry.get("op", "eq") or "eq").strip().lower()
                value = str(filter_entry.get("value", "") or "").strip()
                if not field or not value or op not in {"eq", "ne"}:
                    continue
                filters.append(GigaChatRulePackFilter(field=field, op=op, value=value))

            code = str(entry.get("code", "") or "").strip() or f"RULE_{index + 1}"
            rule_type = str(entry.get("type", "assign_tag") or "assign_tag").strip().lower()
            if rule_type not in {"assign_tag", "reclass_topic"}:
                rule_type = "assign_tag"

            items.append(
                GigaChatRulePack(
                    code=code,
                    description=str(entry.get("description", "") or "").strip(),
                    enabled=bool(entry.get("enabled", True)),
                    type=rule_type,
                    source_fields=cls._split_rule_list(entry.get("source_fields")),
                    keywords=cls._split_keyword_list(entry.get("keywords")),
                    filters=filters,
                    target_tag=str(entry.get("target_tag", "") or "").strip() or None,
                    target_topic=str(entry.get("target_topic", "") or "").strip() or None,
                )
            )
        return items

    @staticmethod
    def _normalize_match_text(value: Any, *, strip: bool = True) -> str:
        text = "" if value is None else str(value)
        if strip:
            text = text.strip()
        text = text.lower()
        text = re.sub(r"\s+", " ", text)
        return text

    @classmethod
    def _keyword_matches_text(cls, keyword: str, text: str) -> bool:
        normalized_keyword = cls._normalize_match_text(keyword, strip=False)
        normalized_text = cls._normalize_match_text(text, strip=False)
        if not normalized_keyword.strip() or not normalized_text:
            return False
        regex_pattern = ".*".join(re.escape(part) for part in normalized_keyword.split("*"))
        regex_pattern = regex_pattern.replace(r"\ ", r"\s+")
        try:
            return re.search(regex_pattern, normalized_text, flags=re.IGNORECASE) is not None
        except re.error:
            return normalized_keyword in normalized_text

    @classmethod
    def _row_filter_matches(cls, row: dict[str, Any], rule_filter: GigaChatRulePackFilter) -> bool:
        field_names = [rule_filter.field, *cls.RULE_FILTER_FIELD_ALIASES.get(rule_filter.field, [])]
        row_value = next((row.get(field_name) for field_name in field_names if field_name in row), None)
        left = cls._normalize_match_text(row_value)
        right = cls._normalize_match_text(rule_filter.value)
        if rule_filter.op == "eq":
            return left == right
        return left != right

    @classmethod
    def _evaluate_rule_hits_for_row(
        cls,
        rule_packs: list[GigaChatRulePack],
        row: dict[str, Any],
        row_index: int,
    ) -> GigaChatRuleEvaluationRow:
        hits: list[GigaChatRuleHit] = []

        for rule_pack in rule_packs:
            if not rule_pack.enabled:
                continue
            if rule_pack.filters and not all(cls._row_filter_matches(row, rule_filter) for rule_filter in rule_pack.filters):
                continue

            matched_keywords: list[str] = []
            matched_fields: list[str] = []
            for keyword in rule_pack.keywords:
                keyword_matched = False
                for field_name in rule_pack.source_fields:
                    if cls._keyword_matches_text(keyword, row.get(field_name)):
                        keyword_matched = True
                        if field_name not in matched_fields:
                            matched_fields.append(field_name)
                if keyword_matched:
                    matched_keywords.append(keyword)

            if not matched_keywords:
                continue

            hits.append(
                GigaChatRuleHit(
                    code=rule_pack.code,
                    description=rule_pack.description,
                    type=rule_pack.type,
                    matched_keywords=matched_keywords,
                    matched_fields=matched_fields,
                    target_tag=rule_pack.target_tag,
                    target_topic=rule_pack.target_topic,
                )
            )

        suggested_tags = list(dict.fromkeys([hit.target_tag for hit in hits if hit.target_tag]))
        suggested_topics = list(dict.fromkeys([hit.target_topic for hit in hits if hit.target_topic]))

        return GigaChatRuleEvaluationRow(
            row_index=row_index,
            hits=hits,
            suggested_tags=[str(item) for item in suggested_tags],
            suggested_topics=[str(item) for item in suggested_topics],
        )

    def evaluate_rule_packs(self, req: GigaChatRuleEvaluationRequest) -> GigaChatRuleEvaluationResponse:
        values = dict(self._effective_values()[0])
        for key, value in req.values.items():
            values[key] = self._coerce_setting_value(key, value)
        rule_packs = self._parse_rule_pack_items(values.get("rule_pack_prompt_notes", ""))
        evaluations = [
            self._evaluate_rule_hits_for_row(rule_packs, row, row_index)
            for row_index, row in enumerate(req.rows)
        ]
        return GigaChatRuleEvaluationResponse(rule_packs=rule_packs, evaluations=evaluations)

    @classmethod
    def _parse_rule_pack_exclusions(cls, raw: Any) -> set[str]:
        if raw is None:
            return set()
        items: list[Any]
        if isinstance(raw, list):
            items = raw
        elif isinstance(raw, dict):
            candidate = raw.get("codes") or raw.get("items") or raw.get("rule_packs")
            items = candidate if isinstance(candidate, list) else [raw]
        else:
            text = str(raw or "").strip()
            if not text:
                return set()
            try:
                parsed = json.loads(text)
            except Exception:
                parsed = cls._split_rule_list(text)
            if isinstance(parsed, list):
                items = parsed
            elif isinstance(parsed, dict):
                candidate = parsed.get("codes") or parsed.get("items") or parsed.get("rule_packs")
                items = candidate if isinstance(candidate, list) else [parsed]
            else:
                items = [parsed]

        codes: set[str] = set()
        for item in items:
            if isinstance(item, dict):
                value = item.get("code") or item.get("name") or item.get("value")
            else:
                value = item
            code = str(value or "").strip()
            if code:
                codes.add(code)
        return codes

    @staticmethod
    def _filter_rule_evaluation_for_prompt(
        evaluation: GigaChatRuleEvaluationRow,
        excluded_codes: set[str],
    ) -> GigaChatRuleEvaluationRow:
        if not excluded_codes:
            return evaluation
        hits = [hit for hit in evaluation.hits if hit.code not in excluded_codes]
        suggested_tags = list(dict.fromkeys([hit.target_tag for hit in hits if hit.target_tag]))
        suggested_topics = list(dict.fromkeys([hit.target_topic for hit in hits if hit.target_topic]))
        return GigaChatRuleEvaluationRow(
            row_index=evaluation.row_index,
            hits=hits,
            suggested_tags=[str(item) for item in suggested_tags],
            suggested_topics=[str(item) for item in suggested_topics],
        )

    @classmethod
    def _format_rule_pack_lines(cls, raw: Any, exclusions_raw: Any = None) -> str:
        rule_packs = cls._parse_rule_pack_items(raw)
        excluded_codes = cls._parse_rule_pack_exclusions(exclusions_raw)
        if excluded_codes:
            rule_packs = [item for item in rule_packs if item.code not in excluded_codes]
        if not rule_packs:
            return "Rule packs: список пока пустой."
        lines = ["Rule packs:"]
        for item in rule_packs:
            if not item.enabled:
                continue
            filter_text = (
                "; ".join(f"{rule_filter.field} {rule_filter.op} {rule_filter.value}" for rule_filter in item.filters)
                if item.filters else
                "без бизнес-фильтров"
            )
            keyword_text = ", ".join(item.keywords) if item.keywords else "без keywords"
            action_text = item.target_tag if item.type == "assign_tag" else item.target_topic
            lines.append(
                f"- {item.code} [{item.type}]: поля={', '.join(item.source_fields) or 'не заданы'}; "
                f"фильтры={filter_text}; keywords={keyword_text}; действие={action_text or 'не задано'}."
            )
            if item.description:
                lines.append(f"  Описание: {item.description}")
        return "\n".join(lines)

    @classmethod
    def _parse_reclassification_rules(cls, values: dict[str, Any]) -> list[dict[str, str]]:
        raw = values.get("reclassification_prompt_notes", "")
        explicit_rules_config = False
        raw_items: list[Any] = []
        if isinstance(raw, list):
            explicit_rules_config = True
            raw_items = raw
        elif isinstance(raw, dict):
            explicit_rules_config = True
            raw_items = [raw]
        else:
            raw_text = str(raw or "").strip()
            if raw_text:
                explicit_rules_config = True
                try:
                    parsed = json.loads(raw_text)
                except Exception:
                    parsed = []
                if isinstance(parsed, list):
                    raw_items = parsed
                elif isinstance(parsed, dict):
                    raw_items = [parsed]

        rules: list[dict[str, Any]] = []
        for item in raw_items:
            if not isinstance(item, dict):
                continue
            name = str(item.get("name") or item.get("title") or item.get("code") or "").strip()
            source_field = str(item.get("source_field") or item.get("sourceField") or "").strip()
            context_fields = cls._split_rule_list(item.get("context_fields") or item.get("contextFields"))
            legacy_context_field = str(item.get("context_field") or item.get("contextField") or "").strip()
            if legacy_context_field and legacy_context_field not in context_fields:
                context_fields.insert(0, legacy_context_field)
            if not source_field and not context_fields:
                continue
            prompt = str(item.get("prompt") or item.get("description") or "").strip()
            rules.append({
                "name": name,
                "source_field": source_field,
                "context_field": context_fields[0] if context_fields else "",
                "context_fields": context_fields,
                "prompt": prompt,
            })

        if rules or explicit_rules_config:
            return rules

        source_field = str(values.get("reclassification_source_field", "") or "").strip()
        context_field = str(values.get("reclassification_context_field", "") or "").strip()
        if not source_field and not context_field:
            return []
        prompt = str(values.get("reclassification_prompt", "") or "").strip()
        return [{
            "name": "",
            "source_field": source_field,
            "context_field": context_field,
            "context_fields": [context_field] if context_field else [],
            "prompt": prompt,
        }]

    def _format_reclassification_lines(self, values: dict[str, Any]) -> str:
        rules = self._reclassification_choice_items(values)
        if not rules:
            return ""
        lines = [
            "Переклассификация: закрытый список допустимых итоговых тем.",
            "- `reclassified_topic` может быть только точным значением `name` из списка ниже или пустой строкой.",
            "- Не придумывай новые темы, не переформулируй названия и не возвращай описание.",
        ]
        for index, rule in enumerate(rules, start=1):
            lines.extend([
                f"{index}. name: {rule['name']}",
                f"   description: {rule['description'] or 'без описания'}",
            ])
        return "\n".join(lines)

    @classmethod
    def _merge_prompt_columns(cls, columns: list[str], values: dict[str, Any]) -> list[str]:
        merged: list[str] = []
        for column in [
            *columns,
            *[
                field
                for rule in cls._parse_reclassification_rules(values)
                for field in [rule["source_field"], *rule.get("context_fields", [])]
            ],
        ]:
            if column and column not in merged:
                merged.append(column)
        return merged

    @classmethod
    def _reclassification_choice_items(cls, values: dict[str, Any]) -> list[dict[str, str]]:
        items: list[dict[str, str]] = []
        seen: set[str] = set()
        for rule in cls._parse_reclassification_rules(values):
            name = str(rule.get("name") or "").strip()
            if not name or name in seen:
                continue
            seen.add(name)
            items.append({
                "name": name,
                "description": str(rule.get("prompt") or "").strip(),
            })
        return items

    @classmethod
    def _reclassification_prompt_columns(cls, values: dict[str, Any], fallback_columns: list[str]) -> list[str]:
        merged: list[str] = []
        for rule in cls._parse_reclassification_rules(values):
            for column in [rule.get("source_field"), *rule.get("context_fields", [])]:
                column = str(column or "").strip()
                if column and column not in merged:
                    merged.append(column)
        if merged:
            return merged
        return list(fallback_columns)

    @classmethod
    def _normalize_reclassification_topic_value(cls, value: Any) -> str:
        text = cls._cell_text(value)
        if not text:
            return ""
        normalized = re.sub(r"\s+", " ", text).casefold()
        if normalized in cls.EMPTY_RECLASSIFICATION_MARKERS:
            return ""
        return text

    @classmethod
    def _sanitize_reclassification_response(cls, response_json: Any, values: dict[str, Any]) -> Any:
        if not isinstance(response_json, dict):
            return response_json
        allowed_names = [item["name"] for item in cls._reclassification_choice_items(values)]
        allowed_by_normalized = {name.strip().casefold(): name for name in allowed_names}
        raw_topic = cls._normalize_reclassification_topic_value(response_json.get("reclassified_topic"))
        sanitized = dict(response_json)
        sanitized["reclassified_topic"] = allowed_by_normalized.get(raw_topic.casefold(), "") if raw_topic else ""
        return sanitized

    @staticmethod
    def _format_row_rule_context(evaluation: GigaChatRuleEvaluationRow) -> str:
        lines: list[str] = []
        if not evaluation.hits:
            lines.append("Rule-based precheck для этой строки не нашел совпадений.")
        else:
            lines.append("Rule-based precheck для этой строки нашел следующие совпадения:")
            for hit in evaluation.hits:
                action = f"tag={hit.target_tag}" if hit.type == "assign_tag" else f"topic={hit.target_topic}"
                lines.append(
                    f"- {hit.code} [{hit.type}] -> {action}; keywords={', '.join(hit.matched_keywords)}; "
                    f"fields={', '.join(hit.matched_fields)}."
                )
            if evaluation.suggested_tags:
                lines.append(f"Предложенные теги по локальным правилам: {', '.join(evaluation.suggested_tags)}")
            if evaluation.suggested_topics:
                lines.append(f"Предложенные переклассификации по локальным правилам: {', '.join(evaluation.suggested_topics)}")
        lines.append(
            "Подтверди или отвергни каждое срабатывание. В confirmed_rule_hits верни подтвержденные коды правил, "
            "в rejected_rule_hits — отклоненные. Если rule-pack предлагает reclass_topic, при подтверждении верни его в reclassified_topic."
        )
        lines.append(
            "Отдельно выполни semantic tag review: локальные совпадения — это подсказка, а не полный результат. "
            "Самостоятельно проверь все разрешенные теги по смыслу в разрешенных source_fields rule-pack'ов: "
            "точные ключи, однокоренные формы, опечатки, искаженные написания, синонимы и косвенное описание ситуации. "
            "Если разрешенный тег подходит без local rule hit, добавь его в model_added_tags. "
            "Если local rule hit ошибочный по контексту или отрицанию, добавь тег в model_rejected_tags и код правила в rejected_rule_hits."
        )
        return "\n".join(lines)

    @staticmethod
    def _placeholder_row(columns: list[str]) -> dict[str, str]:
        if not columns:
            return {"complaint_text": "{{complaint_text}}"}
        return {column: f"{{{{{column}}}}}" for column in columns}

    @classmethod
    def _render_payload_for_row(cls, payload: Any, row: dict[str, Any]) -> Any:
        if isinstance(payload, dict):
            return {key: cls._render_payload_for_row(value, row) for key, value in payload.items()}
        if isinstance(payload, list):
            return [cls._render_payload_for_row(item, row) for item in payload]
        if isinstance(payload, str):
            rendered = payload
            for column, value in row.items():
                rendered = rendered.replace(f"{{{{{column}}}}}", "" if value is None else str(value))
            return rendered
        return payload

    def _compose_final_payload(self, values: dict[str, Any], columns: list[str]) -> dict[str, Any]:
        model = str(values.get("model", self.cfg.llm.model) or self.cfg.llm.model)
        temperature = float(values.get("temperature", getattr(self.cfg.llm, "temperature", 0.2)) or 0.2)
        top_p = float(values.get("top_p", getattr(self.cfg.llm, "top_p", 0.95)) or 0.95)
        max_tokens = int(values.get("max_output_tokens", getattr(self.cfg.llm, "max_output_tokens", 2048)) or 2048)
        system_prompt = str(values.get("system_prompt", self.DEFAULT_SYSTEM_PROMPT) or "").strip() or self.DEFAULT_SYSTEM_PROMPT
        user_prompt_prefix = str(values.get("user_prompt_prefix", getattr(self.cfg.llm, "user_prompt_prefix", "")) or "").strip()
        context_notes = str(values.get("context_notes", getattr(self.cfg.llm, "context_notes", "")) or "").strip()

        rule_pack_rules = self._format_rule_pack_lines(
            values.get("rule_pack_prompt_notes", ""),
            values.get("rule_pack_exclusion_notes", ""),
        )
        reclassification_rules = self._format_reclassification_lines(values)
        prompt_columns = self._merge_prompt_columns(columns, values)
        placeholder_row = self._placeholder_row(prompt_columns)

        prompt_parts: list[str] = []
        if context_notes:
            prompt_parts.append(context_notes)
        if user_prompt_prefix:
            prompt_parts.append(user_prompt_prefix)
        if reclassification_rules:
            prompt_parts.append(reclassification_rules)
        prompt_parts.append(rule_pack_rules)
        prompt_parts.append(
            "Semantic tag review:\n"
            "- Можно назначать только теги из действий активных assign_tag rule packs.\n"
            "- Local rule hits являются подсказками, но не ограничивают результат: проверь разрешенные теги независимо.\n"
            "- Для каждого итогового тега укажи `tag_decisions`: tag, decision (`confirmed`, `added`, `rejected`, `not_applicable`), "
            "source (`local_rule`, `llm_semantic`), match_type (`exact`, `stem`, `typo`, `synonym`, `semantic`, `rejected`, `none`), evidence, reason.\n"
            "- `local_tags` = теги, предложенные локальными правилами. `model_added_tags` = разрешенные теги, добавленные моделью без local hit. "
            "`model_rejected_tags` = локальные теги, отвергнутые моделью. `assigned_tags` = финальные теги после подтверждения/добавления/отклонения.\n"
            "- Если тег добавлен по смыслу, evidence должен быть короткой фразой из строки или точным указанием поля, иначе тег не ставь."
        )
        prompt_parts.append(f"Активные колонки для анализа: {', '.join(prompt_columns) if prompt_columns else 'не выбраны'}")
        prompt_parts.append(
            "Ниже шаблон одной записи. В реальном запросе на место плейсхолдеров будут подставлены значения выбранных колонок:"
        )
        prompt_parts.append(json.dumps(placeholder_row, ensure_ascii=False, indent=2))
        prompt_parts.append(
            "Перед отправкой каждой реальной строки локальный rule-based precheck вычислит совпадения по rule packs и "
            "добавит их в user message. Используй эти совпадения как сильную подсказку, но подтверждай их только если они реально согласуются с текстом строки."
        )

        return {
            "model": model,
            "temperature": temperature,
            "top_p": top_p,
            "max_tokens": max_tokens,
            "messages": [
                {"role": "system", "content": system_prompt},
                {"role": "user", "content": "\n\n".join(prompt_parts)},
            ],
        }

    def _compose_reclassification_payload(self, values: dict[str, Any], columns: list[str]) -> dict[str, Any]:
        model = str(values.get("model", self.cfg.llm.model) or self.cfg.llm.model)
        temperature = float(values.get("temperature", getattr(self.cfg.llm, "temperature", 0.2)) or 0.2)
        top_p = float(values.get("top_p", getattr(self.cfg.llm, "top_p", 0.95)) or 0.95)
        max_tokens = int(values.get("max_output_tokens", getattr(self.cfg.llm, "max_output_tokens", 2048)) or 2048)
        system_prompt = str(values.get("system_prompt", self.DEFAULT_SYSTEM_PROMPT) or "").strip() or self.DEFAULT_SYSTEM_PROMPT
        prompt_columns = self._reclassification_prompt_columns(values, columns)
        placeholder_row = self._placeholder_row(prompt_columns)
        allowed_topics = self._reclassification_choice_items(values)

        prompt_parts = [
            "Задача: проверить, нужно ли заменить исходную тему обращения на одну тему из закрытого списка `allowed_topics`.",
            (
                "Правила ответа:\n"
                "1. Верни только JSON без markdown: {\"reclassified_topic\":\"...\"}.\n"
                "2. Значение `reclassified_topic` должно точно совпадать с одним из `allowed_topics[].name`.\n"
                "3. Сравни исходную тему строки, контекст обращения и описания allowed_topics по смыслу, а не по буквальному тексту.\n"
                "4. Если исходная тема уже означает то же самое, что подходящая тема из allowed_topics, даже если написана свободно, короче или другими словами, верни пустую строку.\n"
                "5. Верни название allowed topic только если контекст явно показывает, что исходная тема ошибочная, слишком общая или относится к другой смысловой теме.\n"
                "6. Если ни одна тема не подходит уверенно или данных недостаточно, верни пустую строку.\n"
                "7. Не возвращай '-', '—', 'нет', null, description и не придумывай новые темы."
            ),
            "allowed_topics:",
            json.dumps(allowed_topics, ensure_ascii=False, indent=2),
            "Данные строки:",
            json.dumps(placeholder_row, ensure_ascii=False, indent=2),
        ]

        return {
            "model": model,
            "temperature": temperature,
            "top_p": top_p,
            "max_tokens": max_tokens,
            "messages": [
                {
                    "role": "system",
                    "content": (
                        f"{system_prompt}\n"
                        "Для этого запроса действует закрытый список: нельзя возвращать значения вне allowed_topics[].name. "
                        "Если текущая тема совпадает с контекстом по смыслу, reclassified_topic должен быть пустой строкой."
                    ),
                },
                {"role": "user", "content": "\n\n".join(prompt_parts)},
            ],
        }

    @staticmethod
    def _payload_to_count_input(payload: dict[str, Any]) -> str:
        messages = payload.get("messages")
        if isinstance(messages, list):
            parts: list[str] = []
            for item in messages:
                if not isinstance(item, dict):
                    continue
                content = item.get("content")
                if content in (None, ""):
                    continue
                parts.append(str(content))
            if parts:
                return "\n\n".join(parts)
        return json.dumps(payload, ensure_ascii=False)

    @classmethod
    def _inject_row_rule_context(cls, payload: dict[str, Any], evaluation: GigaChatRuleEvaluationRow) -> dict[str, Any]:
        messages = payload.get("messages")
        if not isinstance(messages, list):
            return payload

        next_payload = dict(payload)
        next_messages: list[dict[str, Any]] = []
        injected = False
        rule_context = cls._format_row_rule_context(evaluation)

        for message in messages:
            if not isinstance(message, dict):
                next_messages.append(message)
                continue
            next_message = dict(message)
            if next_message.get("role") == "user":
                content = str(next_message.get("content", "") or "").strip()
                next_message["content"] = f"{content}\n\n{rule_context}".strip()
                injected = True
            next_messages.append(next_message)

        if not injected:
            next_messages.append({"role": "user", "content": rule_context})

        next_payload["messages"] = next_messages
        return next_payload

    def _build_sheet_preview(self, sheet_name: str, df: pd.DataFrame) -> GigaChatWorkbookSheetPreview:
        columns = [str(c) for c in df.columns]
        preview_columns = columns[: min(len(columns), 6)]
        preview_df = df[preview_columns].head(5).copy() if preview_columns else df.head(5).copy()
        return GigaChatWorkbookSheetPreview(
            name=sheet_name,
            rows_total=int(len(df)),
            column_count=int(len(columns)),
            columns=columns,
            preview_rows=self._frame_to_rows(preview_df),
        )

    @classmethod
    def _frame_to_rows(cls, df: pd.DataFrame) -> list[dict[str, Any]]:
        out: list[dict[str, Any]] = []
        work = df.copy()
        work.columns = [str(c) for c in work.columns]
        for _, row in work.iterrows():
            out.append({col: cls._normalize_cell(row[col]) for col in work.columns})
        return out

    @staticmethod
    def _normalize_cell(value: Any) -> Any:
        if value is None:
            return None
        if isinstance(value, pd.Series):
            values = [GigaChatLabService._normalize_cell(item) for item in value.tolist()]
            values = [item for item in values if item not in (None, "")]
            return ", ".join(str(item) for item in values) if values else None
        if isinstance(value, (datetime, date)):
            return value.isoformat()
        if isinstance(value, pd.Timestamp):
            return value.isoformat()
        try:
            if pd.isna(value):
                return None
        except Exception:
            pass
        if hasattr(value, "item") and not isinstance(value, (str, bytes, bytearray)):
            try:
                value = value.item()
            except Exception:
                pass
        if isinstance(value, float) and math.isnan(value):
            return None
        if isinstance(value, (bool, int, float, str)):
            return value
        return str(value)

    @classmethod
    def _normalize_export_cell(cls, value: Any) -> Any:
        normalized = cls._normalize_cell(value)
        if isinstance(normalized, str):
            text = normalized.strip()
            if re.fullmatch(r"\d{4}-\d{2}-\d{2}", text):
                try:
                    return datetime.fromisoformat(text).date()
                except Exception:
                    return normalized
            if re.fullmatch(r"\d{4}-\d{2}-\d{2}[ T]\d{2}:\d{2}(:\d{2})?", text):
                try:
                    return datetime.fromisoformat(text.replace(" ", "T"))
                except Exception:
                    return normalized
        return normalized

    @staticmethod
    def _source_export_columns(source_columns: list[str], reserved_columns: list[str]) -> list[tuple[str, str]]:
        used = {str(column) for column in reserved_columns}
        result: list[tuple[str, str]] = []
        for source_column in source_columns:
            source_column = str(source_column).strip()
            if not source_column:
                continue
            base = source_column if source_column not in used else f"Source: {source_column}"
            export_column = base
            suffix = 2
            while export_column in used:
                export_column = f"{base} ({suffix})"
                suffix += 1
            used.add(export_column)
            result.append((source_column, export_column))
        return result

    @staticmethod
    def _excel_safe_sheet_name(value: str) -> str:
        cleaned = re.sub(r"[:\\\\/?*\\[\\]]", "_", str(value or "Разметка")).strip()
        return (cleaned or "Разметка")[:31]

    @staticmethod
    def _sanitize_for_excel(df: pd.DataFrame, max_len: int = 32767) -> pd.DataFrame:
        out = df.copy()
        for col_idx, dtype in enumerate(out.dtypes):
            if not (pd.api.types.is_object_dtype(dtype) or pd.api.types.is_string_dtype(dtype) or str(dtype) == "str"):
                continue
            series = out.iloc[:, col_idx]
            mask = series.map(lambda v: isinstance(v, str))
            if not bool(mask.any()):
                continue
            sanitized = series.loc[mask].astype(str).str.replace(_OPENPYXL_ILLEGAL_RE, "", regex=True).str.slice(0, max_len)
            out.iloc[mask.to_numpy(), col_idx] = sanitized.to_numpy()
        return out
