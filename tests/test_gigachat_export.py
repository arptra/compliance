from __future__ import annotations

from io import BytesIO

import pandas as pd
from openpyxl import load_workbook

from complaints_trends.api.schemas import GigaChatAnnotatedExportRequest, GigaChatAnnotatedExportRow
from complaints_trends.api.services.gigachat_lab_service import GigaChatLabService


def test_annotated_export_renames_source_columns_that_conflict_with_model_columns() -> None:
    service = GigaChatLabService.__new__(GigaChatLabService)
    req = GigaChatAnnotatedExportRequest(
        filename="source.xlsx",
        sheet_name="Sheet1",
        source_columns=["Теги", "Класс", "client_text"],
        rows=[
            GigaChatAnnotatedExportRow(
                row_index=0,
                classification="Жалоба",
                tags=["DRA", "ИПОТЕКА"],
                source_row={
                    "Теги": "исходный тег",
                    "Класс": "исходный класс",
                    "client_text": "Текст клиента",
                },
            ),
        ],
    )

    _, content = service.export_annotated_workbook(req)
    workbook = load_workbook(BytesIO(content))
    sheet = workbook["Sheet1"]
    headers = [cell.value for cell in sheet[1]]

    assert len(headers) == len(set(headers))
    assert "Теги" in headers
    assert "Source: Теги" in headers
    assert "Source: Класс" in headers

    values = {header: sheet.cell(row=2, column=index + 1).value for index, header in enumerate(headers)}
    assert values["Теги"] == "DRA, ИПОТЕКА"
    assert values["Source: Теги"] == "исходный тег"
    assert values["Source: Класс"] == "исходный класс"


def test_sanitize_for_excel_handles_duplicate_dataframe_columns() -> None:
    df = pd.DataFrame([["ok", "bad\x01value"]], columns=["dup", "dup"])

    sanitized = GigaChatLabService._sanitize_for_excel(df)

    assert list(sanitized.columns) == ["dup", "dup"]
    assert sanitized.iloc[0, 0] == "ok"
    assert sanitized.iloc[0, 1] == "badvalue"
